"""Synthetic Forms exports only; no customer participant files are fixtures."""

import csv
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import Workbook
from pydantic import ValidationError

from app.domain.catalog_v4 import SURVEY_VERSION
from app.forms_v4 import (FormsCommitRequestV4, FormsMappingV4, FormsPreviewRequestV4,
                          FormsTargetV4, commit, inspect_columns, preview, purge_case_imports)
from app.input_models_v4 import new_session_v4
from app.maintenance import backup, clean, references, restore
from app.storage import Store


def csv_bytes(rows):
    stream = io.StringIO(newline="")
    csv.writer(stream).writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows, sheet="응답"):
    book = Workbook()
    book.active.title = sheet
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


class FormsImportV4Tests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="kdog-forms-v4-")
        self.addCleanup(temporary.cleanup)
        self.store = Store(Path(temporary.name))
        with self.store.connect(write=True) as db:
            for name, role in (("operator", "operator"), ("reviewer", "reviewer")):
                db.execute("INSERT INTO users(username,password_hash,role) VALUES (?,?,?)", (name, "unused", role))
        self.data = csv_bytes([["반려견", "문항1", "두려움"], ["합성견", "보통", "0"]])
        self.request = FormsPreviewRequestV4(request_id="preview-1", event_id="TEST", filename="synthetic.csv", format="csv",
            mapping=FormsMappingV4(version="mapping-1", source_profile="합성 표준 응답", survey_version=SURVEY_VERSION,
                                  columns={"dog_name": "반려견", "s01": "문항1", "s10": "두려움"}))

    def stage(self, data=None, request=None):
        return preview(self.store, self.data if data is None else data, request or self.request, "operator")

    def apply(self, result, request_id="commit-1", actor="operator"):
        return commit(self.store, FormsCommitRequestV4(request_id=request_id, preview_id=result.preview_id,
                                                      preview_hash=result.preview_hash), actor)

    def manifest(self, case_id):
        with self.store.connect() as db:
            return self.store.manifest(self.store.case(db, case_id))

    def count(self, table="cases"):
        with self.store.connect() as db:
            return db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]

    def audit(self, action, target=None):
        with self.store.connect() as db:
            sql = "SELECT detail_json FROM changes WHERE action=?" + (" AND target=?" if target else "")
            return [json.loads(row[0]) for row in db.execute(sql, (action, target) if target else (action,))]

    def expect_http(self, code, callable_, *args, **kwargs):
        with self.assertRaises(HTTPException) as caught:
            callable_(*args, **kwargs)
        self.assertEqual(caught.exception.status_code, code, caught.exception.detail)

    def linked(self, first, **columns):
        row = first.rows[0]
        request = self.request.model_copy(deep=True)
        request.request_id = "preview-update"
        request.targets = [FormsTargetV4(row_number=2, case_id=row.case_id, session_id=row.session_id,
                                         expected_revision=self.manifest(row.case_id).input_revision)]
        if columns:
            request.mapping.columns = columns
        return request

    def test_ids_are_fixed_in_preview_and_create_and_survey_commit_together(self):
        staged = self.stage()
        self.assertFalse(staged.errors)
        self.assertEqual(self.count(), 0)
        self.assertEqual(staged.rows[0].answers["s10"], 0)
        self.assertEqual(staged.rows[0].registration_status, "partial")
        result = self.apply(staged)
        self.assertEqual((result.rows[0].case_id, result.rows[0].session_id, result.rows[0].participant_id),
                         (staged.rows[0].case_id, staged.rows[0].session_id, staged.rows[0].participant_id))
        manifest = self.manifest(result.rows[0].case_id)
        self.assertEqual(manifest.sessions[0].survey["s01"], 3)
        self.assertIsNone(manifest.sessions[0].survey["s02"])
        self.assertEqual(manifest.consents.analysis_feedback, "unknown")
        self.assertEqual(manifest.consents.stranger_contact, "unknown")

    def test_same_names_and_timestamps_do_not_merge_participants(self):
        data = csv_bytes([["반려견", "문항1", "두려움", "타임스탬프"],
                          ["동명이견", "3", "0", "2026-01-01"], ["동명이견", "3", "0", "2026-01-01"]])
        staged = self.stage(data)
        self.assertFalse(staged.errors)
        self.assertEqual(len({row.participant_id for row in staged.rows}), 2)
        self.apply(staged)
        self.assertEqual(self.count(), 2)

    def test_request_and_file_idempotency_preserve_ids_and_revision(self):
        staged = self.stage()
        self.assertEqual(self.stage(), staged)
        different_request = self.request.model_copy(deep=True)
        different_request.request_id, different_request.filename = "retry-preview", "renamed.csv"
        self.assertEqual(self.stage(request=different_request).preview_id, staged.preview_id)
        result = self.apply(staged)
        self.assertEqual(self.apply(staged), result)
        self.assertEqual(self.apply(staged, "retry-commit"), result)
        self.assertTrue(self.stage(request=different_request).already_committed)
        self.assertEqual(self.manifest(result.rows[0].case_id).input_revision, 1)
        self.assertEqual(self.count(), 1)
        self.assertEqual(len(self.audit("forms.source_row")), 1)

    def test_request_id_different_body_is_rejected_for_preview_and_commit(self):
        staged = self.stage()
        self.expect_http(409, self.stage, self.data + b"\n")
        changed = self.request.model_copy(deep=True)
        changed.filename = "different.csv"
        self.expect_http(409, self.stage, request=changed)
        self.apply(staged)
        self.expect_http(409, commit, self.store, FormsCommitRequestV4(request_id="commit-1", preview_id=staged.preview_id,
                                                                    preview_hash="0" * 64), "operator")

    def test_error_in_later_row_blocks_entire_commit(self):
        staged = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["정상합성", "3", "0"], ["오류합성", "maybe", "0"]]))
        self.assertEqual(staged.errors[0].row_number, 3)
        self.expect_http(422, self.apply, staged)
        self.assertEqual(self.count(), 0)
        self.assertFalse(self.audit("forms.source_row"))

    def test_storage_failure_on_second_row_rolls_back_all_rows(self):
        staged = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["합성1", "3", "0"], ["합성2", "3", "0"]]))
        original = self.store.write_manifest
        calls = []
        def fail_second(manifest):
            calls.append(manifest.case_id)
            if len(calls) == 2:
                raise OSError("synthetic storage failure")
            return original(manifest)
        with patch.object(self.store, "write_manifest", side_effect=fail_second):
            with self.assertRaises(OSError):
                self.apply(staged)
        self.assertEqual(self.count(), 0)
        self.assertFalse(self.audit("forms.source_row"))
        self.assertFalse(self.audit("forms.commit"))
        self.assertEqual(len(self.apply(staged).rows), 2)

    def test_existing_target_shows_before_after_and_adds_revision(self):
        first = self.stage()
        self.apply(first)
        request = self.linked(first, s01="문항1", s10="두려움", s02_reason="결측사유")
        updated = self.stage(csv_bytes([["문항1", "두려움", "결측사유"], ["5", "2", "합성 미응답"]]), request)
        self.assertFalse(updated.errors)
        self.assertEqual(updated.rows[0].previous_answers["s01"], 3)
        self.assertEqual(updated.rows[0].answers["s01"], 5)
        self.assertIn("s02", updated.rows[0].changed_questions)
        result = self.apply(updated, "update-commit")
        self.assertEqual(result.rows[0].input_revision, 2)
        manifest = self.manifest(first.rows[0].case_id)
        self.assertEqual(manifest.sessions[0].survey_blank_reasons, {"s02": "합성 미응답"})
        self.assertEqual(len(manifest.prior_inputs), 1)
        self.assertEqual(len(self.audit("forms.source_row")), 2)

    def test_revision_changed_after_preview_blocks_every_write(self):
        first = self.stage()
        self.apply(first)
        updated = self.stage(self.data, self.linked(first))
        with self.store.connect(write=True) as db:
            row = self.store.case(db, first.rows[0].case_id)
            manifest = self.store.manifest(row)
            manifest.sessions[0].note = "동시 변경"
            self.store.save(db, row, manifest, "operator", "test.change")
        self.expect_http(409, self.apply, updated, "update-commit")
        self.assertEqual(self.manifest(first.rows[0].case_id).input_revision, 2)

    def test_selected_session_changed_after_preview_blocks_commit(self):
        first = self.stage()
        self.apply(first)
        updated = self.stage(self.data, self.linked(first))
        with self.store.connect(write=True) as db:
            row = self.store.case(db, first.rows[0].case_id)
            manifest = self.store.manifest(row)
            manifest.sessions.append(new_session_v4("other-session"))
            manifest.selected_session_id = "other-session"
            self.store.save(db, row, manifest, "operator", "test.session")
        self.expect_http(409, self.apply, updated, "update-commit")

    def test_deleted_participant_blocks_commit_and_idempotent_response(self):
        first = self.stage()
        self.apply(first)
        updated = self.stage(self.data, self.linked(first))
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (first.rows[0].case_id,))
        self.expect_http(403, self.apply, updated, "update-commit")
        self.expect_http(403, self.apply, first)
        self.expect_http(403, self.stage)

    def test_permissions_are_rechecked_even_for_completed_requests(self):
        staged = self.stage()
        self.expect_http(403, self.apply, staged, actor="reviewer")
        self.apply(staged)
        with self.store.connect(write=True) as db:
            db.execute("UPDATE users SET active=0 WHERE username='operator'")
        self.expect_http(403, self.apply, staged)
        self.expect_http(403, self.stage)

    def test_explicit_id_collision_is_shown_without_name_matching(self):
        request = self.request.model_copy(deep=True)
        request.mapping.columns["participant_id"] = "ID"
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "합성견", "3", "0"]])
        first = self.stage(data, request)
        self.apply(first)
        request.request_id = "new-attempt"
        staged = self.stage(data + b"\n", request)
        self.assertEqual(len(staged.errors), 1)
        self.assertIn("명시", staged.errors[0].message)
        self.assertEqual(self.count(), 1)

    def test_duplicate_headers_formulas_and_missing_mapping_headers_fail(self):
        for rows in ([["반려견", "문항1", "문항1"], ["합성", "3", "0"]],
                     [["반려견", "문항1", " 문항1 "], ["합성", "3", "0"]],
                     [["반려견", "문항1", "두려움", "미사용"], ["합성", "3", "0", "=SUM(1,2)"]],
                     [["반려견", "누락"], ["합성", "3"]]):
            with self.subTest(rows=rows):
                self.expect_http(422, self.stage, csv_bytes(rows))

    def test_labels_are_explicit_and_ambiguous_or_wrong_scale_mappings_fail(self):
        request = self.request.model_copy(deep=True)
        request.mapping.answer_labels = {"s01": {"예시 응답": 4}}
        staged = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["합성", "예시 응답", "0"]]), request)
        self.assertEqual(staged.rows[0].answers["s01"], 4)
        for labels in ({"3": 5}, {"예시": 0}):
            request.request_id += "x"
            request.mapping.answer_labels = {"s01": labels}
            self.expect_http(422, self.stage, request=request)
        with self.assertRaises(ValidationError):
            FormsMappingV4.model_validate({**request.mapping.model_dump(), "answer_labels": {"s01": {"예시": 1, " 예시 ": 2}}})

    def test_xlsx_sheet_labels_and_raw_integer_zero_are_supported(self):
        request = self.request.model_copy(deep=True)
        request.format, request.filename, request.mapping.sheet = "xlsx", "synthetic.xlsx", "응답"
        data = xlsx_bytes([["반려견", "문항1", "두려움"], ["합성견", 3, 0]])
        staged = self.stage(data, request)
        self.assertEqual(staged.source_sheet, "응답")
        self.assertEqual(staged.rows[0].raw_values["문항1"], 3)
        self.assertEqual(staged.rows[0].answers["s10"], 0)
        self.apply(staged)

    def test_column_inspection_shares_layout_header_and_formula_checks(self):
        inspected = inspect_columns(csv_bytes([["반려견", "합성견"], ["문항1", "3"], ["두려움", "0"]]),
                                    "csv", layout="transposed")
        self.assertEqual(inspected["columns"], [{"key": x, "label": x} for x in ("반려견", "문항1", "두려움")])
        self.assertEqual(inspected["sheets"], [])
        self.assertIsNone(inspected["selected_sheet"])
        inspected = inspect_columns(xlsx_bytes([["반려견", "문항1"], ["합성견", 3]]), "xlsx", "응답")
        self.assertEqual(inspected["sheets"], ["응답"])
        self.assertEqual(inspected["selected_sheet"], "응답")
        self.expect_http(422, inspect_columns, csv_bytes([["반려견", "반려견"]]), "csv")
        self.expect_http(422, inspect_columns, xlsx_bytes([["문항1"], ["=1+2"]]), "xlsx")

    def test_xlsx_formula_and_boolean_answer_are_rejected(self):
        request = self.request.model_copy(deep=True)
        request.format = "xlsx"
        self.expect_http(422, self.stage, xlsx_bytes([["반려견", "문항1", "두려움"], ["합성", "=1+2", 0]]), request)
        staged = self.stage(xlsx_bytes([["반려견", "문항1", "두려움"], ["합성", True, 0]]), request)
        self.assertEqual(staged.errors[0].row_number, 2)
        self.expect_http(422, self.apply, staged)

    def test_unsupported_scale_is_reference_only_and_never_converted(self):
        request = self.request.model_copy(deep=True)
        request.mapping.survey_version = "historical-comfort-1-to-5"
        request.mapping.historical_reference_only = True
        request.mapping.source_profile = "합성 과거 전치표"
        request.mapping.layout = "transposed"
        staged = self.stage(csv_bytes([["반려견", "과거합성"], ["문항1", "5"], ["두려움", "1"]]), request)
        self.assertFalse(staged.errors)
        self.assertEqual(staged.rows[0].registration_status, "reference_only")
        self.assertTrue(all(value is None for value in staged.rows[0].answers.values()))
        result = self.apply(staged)
        self.assertEqual(result.rows[0].registration_status, "reference_only")
        self.assertTrue(all(value is None for value in self.manifest(result.rows[0].case_id).sessions[0].survey.values()))
        link = self.audit("forms.source_row")[0]
        capsule = json.loads(self.store.path(link["ref"]).read_bytes())
        self.assertEqual(capsule["raw_values"]["두려움"], "1")
        self.assertEqual(capsule["mapping"]["survey_version"], "historical-comfort-1-to-5")

    def test_unsupported_scale_requires_reference_only_flag(self):
        with self.assertRaises(ValidationError):
            FormsMappingV4.model_validate({**self.request.mapping.model_dump(), "survey_version": "unknown-edition"})

    def test_missing_target_row_or_duplicate_target_is_rejected(self):
        first = self.stage()
        self.apply(first)
        request = self.linked(first)
        request.targets[0].row_number = 3
        staged = self.stage(self.data, request)
        self.assertTrue(any(issue.row_number == 3 for issue in staged.errors))
        self.expect_http(422, self.apply, staged, "update-commit")
        with self.assertRaises(ValidationError):
            FormsPreviewRequestV4.model_validate({**request.model_dump(), "targets": [request.targets[0].model_dump()] * 2})

    def test_source_bytes_hash_mapping_and_rows_are_reachable_for_backup(self):
        staged = self.stage()
        self.assertEqual(staged.source_hash, hashlib.sha256(self.data).hexdigest())
        record_link = self.audit("forms.preview")[0]
        record = json.loads(self.store.path(record_link["ref"]).read_bytes())
        self.assertEqual(self.store.path(record["source"]["ref"]).read_bytes(), self.data)
        self.assertEqual(record["request"]["mapping"]["version"], "mapping-1")
        self.assertEqual(record["original_rows"][1][0], "합성견")
        self.apply(staged)
        with self.store.connect() as db:
            refs = references(self.store, db)
        self.assertIn(record_link["ref"], refs)
        self.assertIn(record["source"]["ref"], refs)
        self.assertIn(self.audit("forms.source_row")[0]["ref"], refs)

    def test_tampered_original_blocks_commit(self):
        staged = self.stage()
        record = json.loads(self.store.path(self.audit("forms.preview")[0]["ref"]).read_bytes())
        self.store.path(record["source"]["ref"]).write_bytes(b"corrupt synthetic file")
        self.expect_http(409, self.apply, staged)
        self.assertEqual(self.count(), 0)

    def test_tampered_preview_blocks_commit(self):
        staged = self.stage()
        self.store.path(self.audit("forms.preview")[0]["ref"]).write_bytes(b"{}")
        self.expect_http(409, self.apply, staged)

    def test_partial_deletion_purges_shared_source_but_preserves_other_row(self):
        data = csv_bytes([["반려견", "문항1", "두려움"], ["삭제합성", "3", "0"], ["보존합성", "4", "1"]])
        staged = self.stage(data)
        self.apply(staged)
        record_link = self.audit("forms.preview")[0]
        record = json.loads(self.store.path(record_link["ref"]).read_bytes())
        deleted, retained = staged.rows
        kept_link = self.audit("forms.source_row", retained.case_id)[0]
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (deleted.case_id,))
            purge_case_imports(self.store, db, deleted.case_id)
        clean(self.store, purge_deleted=True)
        self.assertFalse(self.store.path(record_link["ref"]).exists())
        self.assertFalse(self.store.path(record["source"]["ref"]).exists())
        self.assertTrue(self.store.path(kept_link["ref"]).exists())
        self.assertEqual(self.manifest(retained.case_id).sessions[0].survey["s01"], 4)
        self.expect_http(409, self.apply, staged)
        request = self.request.model_copy(deep=True)
        request.request_id = "reimport"
        request.mapping.version = "mapping-after-deletion"
        self.expect_http(409, self.stage, data, request)
        tombstones = self.audit("forms.purged")
        self.assertTrue(tombstones)
        self.assertTrue(all(set(value) == {"source_hash"} for value in tombstones))

    def test_error_only_previews_are_purged_by_explicit_target_or_stable_identity(self):
        first = self.stage()
        self.apply(first)
        linked = self.linked(first)
        invalid = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["오류합성", "invalid", "0"]]), linked)
        self.assertFalse(invalid.rows)
        self.assertTrue(invalid.errors)
        collision = self.request.model_copy(deep=True)
        collision.request_id = "collision"
        collision.mapping.columns["participant_id"] = "ID"
        staged = self.stage(csv_bytes([["ID", "반려견", "문항1", "두려움"],
                                      [first.rows[0].participant_id, "중복합성", "3", "0"]]), collision)
        self.assertFalse(staged.rows)
        links = self.audit("forms.preview")
        sources = [json.loads(self.store.path(link["ref"]).read_bytes())["source"] for link in links]
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (first.rows[0].case_id,))
        self.expect_http(403, self.stage, csv_bytes([["반려견", "문항1", "두려움"], ["오류합성", "invalid", "0"]]), linked)
        self.expect_http(403, self.apply, invalid, "commit-error-after-deletion")
        with self.store.connect(write=True) as db:
            purge_case_imports(self.store, db, first.rows[0].case_id)
        clean(self.store, purge_deleted=True)
        self.assertFalse(self.audit("forms.preview"))
        self.assertTrue(all(not self.store.path(link["ref"]).exists() for link in links + sources))

    def test_already_deleted_target_cannot_leave_a_new_preview_blob(self):
        first = self.stage()
        self.apply(first)
        request = self.linked(first)
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (first.rows[0].case_id,))
        initial = len(self.audit("forms.preview"))
        self.expect_http(403, self.stage, self.data, request)
        self.assertEqual(len(self.audit("forms.preview")), initial)

    def test_failed_new_identity_then_corrected_registration_purges_both_sources(self):
        request = self.request.model_copy(deep=True)
        request.mapping.columns["participant_id"] = "ID"
        failed = self.stage(csv_bytes([["ID", "반려견", "문항1", "두려움"], ["P001", "합성견", "bad", "0"]]), request)
        self.assertFalse(failed.rows)
        self.assertTrue(failed.errors)
        request.request_id = "corrected"
        good = self.stage(csv_bytes([["ID", "반려견", "문항1", "두려움"], ["P001", "합성견", "3", "0"]]), request)
        self.apply(good)
        links = self.audit("forms.preview")
        sources = [json.loads(self.store.path(link["ref"]).read_bytes())["source"] for link in links]
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (good.rows[0].case_id,))
        request.request_id = "failed-retry"
        self.expect_http(403, self.stage,
                         csv_bytes([["ID", "반려견", "문항1", "두려움"], ["P001", "합성견", "bad", "0"]]), request)
        with self.store.connect(write=True) as db:
            purge_case_imports(self.store, db, good.rows[0].case_id)
        clean(self.store, purge_deleted=True)
        self.assertFalse(self.audit("forms.preview"))
        self.assertTrue(all(not self.store.path(link["ref"]).exists() for link in links + sources))

    def test_deletion_purges_other_sheet_preview_that_references_same_workbook_bytes(self):
        book = Workbook()
        book.active.title = "첫시트"
        book.create_sheet("둘째시트")
        for index, sheet in enumerate(book.worksheets, 1):
            sheet.append(["ID", "반려견", "문항1", "두려움"])
            sheet.append([f"P00{index}", f"합성{index}", 3, 0])
        out = io.BytesIO()
        book.save(out)
        book.close()
        request = self.request.model_copy(deep=True)
        request.mapping.columns["participant_id"] = "ID"
        request.format = "xlsx"
        request.mapping.sheet = "첫시트"
        first = self.stage(out.getvalue(), request)
        self.apply(first)
        request.request_id = "second-sheet"
        request.mapping.sheet = "둘째시트"
        second = self.stage(out.getvalue(), request)
        self.apply(second, "commit-second")
        links = self.audit("forms.preview")
        sources = [json.loads(self.store.path(link["ref"]).read_bytes())["source"] for link in links]
        kept_link = self.audit("forms.source_row", second.rows[0].case_id)[0]
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (first.rows[0].case_id,))
            purge_case_imports(self.store, db, first.rows[0].case_id)
        clean(self.store, purge_deleted=True)
        self.assertFalse(self.audit("forms.preview"))
        self.assertTrue(all(not self.store.path(link["ref"]).exists() for link in links + sources))
        self.assertTrue(self.store.path(kept_link["ref"]).exists())
        self.assertEqual(self.manifest(second.rows[0].case_id).input_revision, 1)
        self.expect_http(409, self.apply, second, "commit-second")

    def restore_preview_before_case_existed(self, *, failed_first):
        temporary = tempfile.TemporaryDirectory(prefix="kdog-forms-backup-v4-")
        self.addCleanup(temporary.cleanup)
        destination = Path(temporary.name)
        request = self.request.model_copy(deep=True)
        if failed_first:
            request.mapping.columns["participant_id"] = "ID"
            bad = self.stage(csv_bytes([["ID", "반려견", "문항1", "두려움"], ["P001", "합성", "bad", "0"]]), request)
            self.assertTrue(bad.errors)
        else:
            staged = self.stage()
        backup(self.store, destination / "snapshot", "operator")
        if failed_first:
            request.request_id = "corrected"
            staged = self.stage(csv_bytes([["ID", "반려견", "문항1", "두려움"], ["P001", "합성", "3", "0"]]), request)
        self.apply(staged)
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (staged.rows[0].case_id,))
        clean(self.store, purge_deleted=True)
        restore(self.store, destination / "snapshot", destination / "restored")
        restored = Store(destination / "restored")
        with restored.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM cases").fetchone()[0], 0)
            self.assertGreater(db.execute("SELECT COUNT(*) FROM changes WHERE action='deletion.record'").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM changes WHERE action='forms.preview'").fetchone()[0], 0)
            self.assertFalse(any("forms-" in ref for ref in references(restored, db)))
        self.assertFalse(list(restored.path("inputs").glob("forms-*")))

    def test_restore_deletion_ledger_purges_error_preview_before_case_existed(self):
        self.restore_preview_before_case_existed(failed_first=True)

    def test_restore_deletion_ledger_purges_generated_identity_preview_before_case_existed(self):
        self.restore_preview_before_case_existed(failed_first=False)


    # --- U03: 정확한 ID 일치 갱신과 정상 행만 등록 ---
    def id_request(self, request_id="id-preview", **changes):
        request = self.request.model_copy(deep=True)
        request.request_id = request_id
        request.mapping.columns["participant_id"] = "ID"
        for key, value in changes.items():
            setattr(request, key, value)
        return request

    def register_existing(self, participant_id="0001", dog="기존견"):
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], [participant_id, dog, "3", "0"]])
        return self.apply(self.stage(data, self.id_request(f"first-{participant_id}")), f"first-commit-{participant_id}")

    def test_exact_id_match_updates_the_existing_participant_only_when_requested(self):
        done = self.register_existing()
        case_id = done.rows[0].case_id
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "다른이름", "5", "2"]])
        default = self.stage(data, self.id_request("default-match"))
        self.assertIn("이미 등록된 ID", default.errors[0].message)
        staged = self.stage(data, self.id_request("matched", match_existing_ids=True))
        self.assertFalse(staged.errors)
        row = staged.rows[0]
        self.assertEqual((row.action, row.matched_by, row.case_id, row.dog_name), ("update", "id", case_id, "기존견"))
        self.assertEqual(row.expected_revision, 1)
        self.assertEqual((row.previous_answers["s01"], row.answers["s01"]), (3, 5))
        self.assertNotEqual(staged.preview_id, default.preview_id)
        result = self.apply(staged, "matched-commit")
        self.assertEqual(result.rows[0].input_revision, 2)
        manifest = self.manifest(case_id)
        self.assertEqual(manifest.sessions[0].survey["s01"], 5)
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT dog_name FROM cases WHERE case_id=?", (case_id,)).fetchone()[0], "기존견")
        self.assertEqual(self.count(), 1)

    def test_same_name_with_a_different_id_is_a_new_participant(self):
        self.register_existing("0001", "같은이름")
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0002", "같은이름", "5", "2"]])
        staged = self.stage(data, self.id_request("other-id", match_existing_ids=True))
        self.assertEqual((staged.rows[0].action, staged.rows[0].matched_by), ("create", None))
        self.apply(staged, "other-commit")
        self.assertEqual(self.count(), 2)

    def test_matched_update_is_blocked_by_a_revision_change_or_deletion_after_preview(self):
        done = self.register_existing()
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "무관", "5", "2"]])
        staged = self.stage(data, self.id_request("revision", match_existing_ids=True))
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET input_revision=input_revision+1 WHERE case_id=?", (done.rows[0].case_id,))
        self.expect_http(409, self.apply, staged, "revision-commit")
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (done.rows[0].case_id,))
        self.expect_http(403, self.apply, staged, "deleted-commit")

    def test_deleted_participant_is_still_refused_when_matching_is_requested(self):
        done = self.register_existing()
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?", (done.rows[0].case_id,))
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "무관", "5", "2"]])
        self.expect_http(403, self.stage, data, self.id_request("deleted", match_existing_ids=True))

    def test_explicit_target_wins_and_two_rows_cannot_claim_one_participant(self):
        first = self.register_existing("0001", "첫째")
        self.register_existing("0002", "둘째")
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0002", "무관", "5", "2"]])
        request = self.id_request("explicit-wins", match_existing_ids=True)
        request.targets = [FormsTargetV4(row_number=2, case_id=first.rows[0].case_id, session_id=first.rows[0].session_id, expected_revision=1)]
        staged = self.stage(data, request)
        self.assertIn("명시적으로 연결한 대상과 다릅니다", staged.errors[0].message)
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["", "첫째", "4", "1"], ["0001", "중복", "5", "2"]])
        request = self.id_request("double-claim", match_existing_ids=True)
        request.targets = [FormsTargetV4(row_number=2, case_id=first.rows[0].case_id, session_id=first.rows[0].session_id, expected_revision=1)]
        staged = self.stage(data, request)
        self.assertEqual([row.row_number for row in staged.rows], [2])
        self.assertEqual(staged.errors[0].row_number, 3)
        self.assertIn("둘 이상", staged.errors[0].message)

    def test_numeric_or_blank_ids_are_never_matched(self):
        self.register_existing("0001")
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["", "빈ID견", "5", "2"], ["1", "짧은ID견", "5", "2"]])
        staged = self.stage(data, self.id_request("blank-numeric", match_existing_ids=True))
        self.assertEqual([(row.action, row.matched_by) for row in staged.rows], [("create", None), ("create", None)])
        self.assertFalse(staged.errors)

    def test_error_rows_keep_their_original_values_for_a_correction_file(self):
        staged = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["정상합성", "3", "0"], ["오류합성", "maybe", "0"]]))
        self.assertEqual(staged.errors[0].raw_values, {"반려견": "오류합성", "문항1": "maybe", "두려움": "0"})

    def test_partial_commit_registers_only_normal_rows_once_and_reports_exclusions(self):
        staged = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["정상합성", "3", "0"], ["오류합성", "maybe", "0"]]))
        request = FormsCommitRequestV4(request_id="partial-1", preview_id=staged.preview_id, preview_hash=staged.preview_hash, accept_partial=True)
        self.expect_http(422, self.apply, staged)
        self.assertEqual(self.count(), 0)
        result = commit(self.store, request, "operator")
        self.assertEqual([row.row_number for row in result.rows], [2])
        # 영수증에는 참가자의 원본 칸 값을 남기지 않는다(원행은 미리보기 기록에만 있다).
        self.assertEqual([(item.row_number, item.raw_values) for item in result.excluded], [(3, None)])
        self.assertIn("문항1", staged.errors[0].raw_values)
        self.assertEqual(self.count(), 1)
        self.assertEqual(commit(self.store, request, "operator"), result)
        retry = request.model_copy(update={"request_id": "partial-retry"})
        self.assertEqual(commit(self.store, retry, "operator"), result)
        self.assertEqual(self.count(), 1)
        self.assertEqual(len(self.audit("forms.source_row")), 1)
        self.expect_http(409, commit, self.store, request.model_copy(update={"accept_partial": False}), "operator")

    def test_partial_commit_needs_a_normal_row_and_rolls_back_on_failure(self):
        bad = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["오류합성", "maybe", "0"]]))
        request = FormsCommitRequestV4(request_id="partial-none", preview_id=bad.preview_id, preview_hash=bad.preview_hash, accept_partial=True)
        self.expect_http(422, commit, self.store, request, "operator")
        second = self.request.model_copy(deep=True)
        second.request_id = "partial-second"
        staged = self.stage(csv_bytes([["반려견", "문항1", "두려움"], ["정상1", "3", "0"], ["정상2", "3", "0"], ["오류합성", "maybe", "0"]]), second)
        import app.forms_v4 as module
        original, calls = module._create_case, []

        def failing(*args, **kwargs):
            calls.append(1)
            if len(calls) == 2:
                raise HTTPException(409, "synthetic failure on the second row")
            return original(*args, **kwargs)

        request = FormsCommitRequestV4(request_id="partial-fail", preview_id=staged.preview_id, preview_hash=staged.preview_hash, accept_partial=True)
        with patch.object(module, "_create_case", failing):
            self.expect_http(409, commit, self.store, request, "operator")
        self.assertEqual(self.count(), 0)
        self.assertFalse(self.audit("forms.source_row"))

    def test_new_rows_without_a_dog_name_get_a_clear_error_but_existing_ids_do_not_need_one(self):
        self.register_existing("0001")
        request = self.id_request("no-name", match_existing_ids=True)
        del request.mapping.columns["dog_name"]
        data = csv_bytes([["ID", "문항1", "두려움"], ["0001", "5", "2"], ["0002", "4", "1"]])
        staged = self.stage(data, request)
        self.assertEqual([(row.participant_id, row.action) for row in staged.rows], [("0001", "update")])
        self.assertEqual(staged.errors[0].row_number, 3)
        self.assertIn("반려견 이름이 없는 행은 새 참가자로 등록할 수 없습니다", staged.errors[0].message)

    def test_partial_receipt_with_hostile_headers_does_not_break_backup_clean_or_restore(self):
        request = self.id_request("hostile")
        request.mapping.columns = {"participant_id": "ID", "dog_name": "반려견", "s01": "문항1"}
        data = csv_bytes([["ID", "반려견", "문항1", "ref", "storage_ref_json"], ["H1", "정상", "3", "", ""], ["H2", "오류", "maybe", "abc", "{not json"]])
        staged = self.stage(data, request)
        self.assertEqual(len(staged.errors), 1)
        commit(self.store, FormsCommitRequestV4(request_id="hostile-commit", preview_id=staged.preview_id,
                                                preview_hash=staged.preview_hash, accept_partial=True), "operator")
        with self.store.connect() as db:
            self.assertTrue(references(self.store, db))
        backup_path = self.store.root.parent / (self.store.root.name + "-hostile-backup")
        restored = self.store.root.parent / (self.store.root.name + "-hostile-restored")
        import shutil
        self.addCleanup(shutil.rmtree, backup_path, True)
        self.addCleanup(shutil.rmtree, restored, True)
        backup(self.store, backup_path, "operator")
        restore(self.store, backup_path, restored)
        clean(self.store, purge_deleted=True)

    def test_a_stale_matched_preview_is_rebuilt_instead_of_reused(self):
        done = self.register_existing("0001")
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "무관", "5", "2"]])
        first = self.stage(data, self.id_request("stale-1", match_existing_ids=True))
        other = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "무관", "1", "1"]])
        self.apply(self.stage(other, self.id_request("stale-other", match_existing_ids=True)), "stale-other-commit")
        self.assertEqual(self.manifest(done.rows[0].case_id).input_revision, 2)
        self.expect_http(409, self.apply, first, "stale-first-commit")
        fresh = self.stage(data, self.id_request("stale-2", match_existing_ids=True))
        self.assertNotEqual(fresh.preview_id, first.preview_id)
        self.assertEqual(fresh.rows[0].expected_revision, 2)
        self.assertEqual(self.apply(fresh, "stale-fresh-commit").rows[0].input_revision, 3)

    def test_an_id_registered_after_preview_turns_the_next_preview_into_an_update(self):
        data = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0009", "새견", "5", "2"]])
        first = self.stage(data, self.id_request("late-1", match_existing_ids=True))
        self.assertEqual(first.rows[0].action, "create")
        self.register_existing("0009", "먼저접수견")
        self.expect_http(409, self.apply, first, "late-first-commit")
        fresh = self.stage(data, self.id_request("late-2", match_existing_ids=True))
        self.assertEqual((fresh.rows[0].action, fresh.rows[0].matched_by, fresh.rows[0].dog_name), ("update", "id", "먼저접수견"))
        self.assertEqual(fresh.rows[0].file_dog_name, "새견")

    def test_unchanged_matched_rows_keep_the_revision_and_blank_or_unmapped_answers_are_not_erased(self):
        done = self.register_existing("0001")
        case_id = done.rows[0].case_id
        same = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "기존견", "3", "0"]])
        staged = self.stage(same, self.id_request("noop", match_existing_ids=True))
        self.assertEqual(staged.rows[0].changed_questions, [])
        self.assertIsNone(staged.rows[0].file_dog_name)
        result = self.apply(staged, "noop-commit")
        self.assertEqual(result.rows[0].input_revision, 1)
        self.assertEqual(self.manifest(case_id).input_revision, 1)
        self.assertEqual(len(self.audit("forms.source_row")), 2)
        # 명단만 있는 파일과 파일의 빈 칸은 기존 응답을 지우지 않는다.
        roster = self.id_request("roster", match_existing_ids=True)
        roster.mapping.columns = {"participant_id": "ID", "dog_name": "반려견"}
        staged = self.stage(csv_bytes([["ID", "반려견"], ["0001", "기존견"]]), roster)
        self.assertEqual((staged.rows[0].changed_questions, staged.rows[0].answers["s01"], staged.rows[0].answers["s10"]), ([], 3, 0))
        blank = csv_bytes([["ID", "반려견", "문항1", "두려움"], ["0001", "기존견", "", "4"]])
        staged = self.stage(blank, self.id_request("blank", match_existing_ids=True))
        self.assertEqual((staged.rows[0].answers["s01"], staged.rows[0].answers["s10"], staged.rows[0].changed_questions), (3, 4, ["s10"]))

    def test_a_blank_with_a_written_reason_may_clear_an_answer_for_matched_rows(self):
        done = self.register_existing("0001")
        request = self.id_request("reason", match_existing_ids=True)
        request.mapping.columns["s01_reason"] = "사유"
        staged = self.stage(csv_bytes([["ID", "반려견", "문항1", "사유", "두려움"], ["0001", "기존견", "", "보호자가 응답하지 않음", "0"]]), request)
        self.assertFalse(staged.errors)
        self.assertEqual((staged.rows[0].answers["s01"], staged.rows[0].blank_reasons), (None, {"s01": "보호자가 응답하지 않음"}))
        self.apply(staged, "reason-commit")
        session = self.manifest(done.rows[0].case_id).sessions[0]
        self.assertIsNone(session.survey["s01"])
        self.assertEqual(session.survey_blank_reasons, {"s01": "보호자가 응답하지 않음"})


if __name__ == "__main__":
    unittest.main()
