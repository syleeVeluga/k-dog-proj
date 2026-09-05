import hashlib
import json
import re
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from app.domain.contracts import BehaviorCatalog, SurveyCatalog

ROOT = Path(__file__).resolve().parents[2]


class SourceCatalogTests(unittest.TestCase):
    def setUp(self):
        self.behavior = BehaviorCatalog.model_validate_json((ROOT / "resources/catalogs/behavior-v1.json").read_bytes())
        self.survey = SurveyCatalog.model_validate_json((ROOT / "resources/catalogs/survey-v1.json").read_bytes())

    def test_source_hashes_and_counts(self):
        for catalog, count in ((self.behavior, 55), (self.survey, 30)):
            with self.subTest(catalog=catalog.source_filename):
                source = ROOT / "docs" / catalog.source_filename
                self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), catalog.source_sha256)
                self.assertEqual(len(catalog.items), count)
        self.assertEqual(Counter(item.domain for item in self.behavior.items),
                         {"EDU": 7, "SOC_P": 1, "SOC_D": 1, "SOC_E": 9, "ATT": 14, "CON": 12, "TRN": 6, None: 5})
        self.assertEqual(Counter(item.source_layer for item in self.survey.items), {"원": 16, "C": 8, "영": 6})

    def test_every_behavior_text_and_nonblank_option_matches_excel(self):
        book = load_workbook(ROOT / "docs" / self.behavior.source_filename, read_only=True, data_only=False)
        try:
            for item in self.behavior.items:
                with self.subTest(item=item.item_id):
                    sheet = book[item.source_sheet]
                    row = item.source_row
                    self.assertEqual(item.text, sheet[f"B{row}"].value)
                    self.assertEqual(item.segment, sheet[f"A{row}"].value)
                    original = {f"{column}{row}": sheet[f"{column}{row}"].value for column in "CDEFG"
                                if sheet[f"{column}{row}"].value is not None}
                    self.assertEqual({option.source_cell: option.text for option in item.options}, original)
                    for option in item.options:
                        self.assertEqual(option.score, "CDEFG".index(option.source_cell[0]) + 1)
                        expected_direction = "A" if "[A]" in option.text else "B" if "[B]" in option.text else None
                        self.assertEqual(option.direction, expected_direction)
                        self.assertEqual(option.option_id, f"{item.item_id}:S{int(option.score)}")
        finally:
            book.close()

    def test_behavior_ids_and_locations_match_prd_appendix(self):
        prd = (ROOT / "docs/K-DOG_PRD_v0.4_20260905.md").read_text(encoding="utf-8")
        rows = re.findall(r"^\| ((?:BS|DOG|OWN)-\d{2}) \| `([^`]+)` (\d+)행 \| ([^|]+) \| ([^|]+) \| ([^|]+) \|$", prd, re.M)
        self.assertEqual(len(rows), 55)
        actual = {item.item_id: item for item in self.behavior.items}
        for item_id, sheet, row, segment, text, domain in rows:
            item = actual[item_id]
            with self.subTest(item=item_id):
                self.assertEqual((item.source_sheet, item.source_row, item.segment, item.text, item.domain),
                                 (sheet, int(row), segment.strip(), text.strip(), None if domain.strip() == "집계 제외" else domain.strip()))

    def test_every_survey_question_and_scoring_note_matches_excel(self):
        book = load_workbook(ROOT / "docs" / self.survey.source_filename, read_only=True, data_only=False)
        try:
            sheet = book["데이터입력"]
            self.assertEqual(self.survey.response_instructions, sheet["A2"].value)
            for item in self.survey.items:
                row = int(item.item_id[1:]) + 5
                with self.subTest(item=item.item_id):
                    self.assertEqual(item.source_row, row)
                    self.assertEqual(item.source_sheet, sheet.title)
                    self.assertEqual((item.domain_label, item.source_layer, item.text, item.scoring_note),
                                     tuple(sheet[f"{column}{row}"].value for column in "BCDE"))
            self.assertEqual(self.survey.items[22].scoring_note, "※중간이 정상")
            self.assertEqual({item.item_id for item in self.survey.items if item.scoring_note == "역(R)"},
                             {"q22", "q26", "q27", "q28", "q29", "q30"})
        finally:
            book.close()

    def test_pending_rules_do_not_activate_draft_calculations(self):
        rules = json.loads((ROOT / "resources/rules/pending-v1.json").read_text(encoding="utf-8"))
        self.assertEqual(len(rules["rules"]), 4)
        self.assertTrue(all(rule["status"].endswith("pending") for rule in rules["rules"]))
        self.assertIn("q23", rules["rules"][0]["affected_items"])


if __name__ == "__main__":
    unittest.main()
