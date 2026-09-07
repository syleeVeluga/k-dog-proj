"""Explicit synthetic T15 harness. Real media/DB/exports, injected AI only; no keys."""

import argparse
import csv
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import platform
import statistics
import subprocess
import tempfile
import time
from unittest.mock import patch
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import create_app
from app.auth import create_user
from app.input_models import UserCreate
from app.maintenance import backup, restore
from app.storage import Store
from app.usage import summarize
from app.worker import Worker
from tests.evaluation_fixtures import FakeEvaluator
from tests.report_fixtures import FakeReporter
from tests.test_observation import FakeObserver


def checked(response, status=200):
    if response.status_code != status:
        raise RuntimeError(f"pilot request failed: {response.status_code} {response.text[:300]}")
    return response.json()


def scenario(participants=20, cameras=2, seconds=180):
    if min(participants, cameras, seconds) < 1:
        raise ValueError("시나리오 수치는 양수여야 합니다.")
    with tempfile.TemporaryDirectory(prefix="kdog-m6-pilot-") as temporary, patch.dict(os.environ, {
        "KDOG_GEMINI_MODEL": "gemini-test-only", "GEMINI_API_KEY": "synthetic-test-only",
    }):
        root = Path(temporary)
        videos = [root / f"synthetic-{camera}.mp4" for camera in range(cameras)]
        for camera, video in enumerate(videos):
            subprocess.run(["ffmpeg", "-v", "error", "-nostdin", "-f", "lavfi", "-i", "color=c=gray:s=640x360:r=10",
                "-f", "lavfi", "-i", f"sine=frequency={440 + camera * 100}:sample_rate=16000", "-t", str(seconds),
                "-c:v", "libx264", "-preset", "ultrafast", "-pix_fmt", "yuv420p", "-c:a", "aac", str(video)],
                check=True, capture_output=True, timeout=180, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        app = create_app(root / "data")
        store = app.state.store
        with store.connect(write=True) as db:
            create_user(db, UserCreate(username="pilot", role="admin", password="Synthetic-pilot-only-42"))
        observer, evaluator, reporter = FakeObserver(), FakeEvaluator(), FakeReporter()
        worker = Worker(store, observer=observer, evaluator=evaluator, reporter=reporter)
        durations, cases = [], []
        expected_ids = {f"{index + 1:04}" for index in range(participants)}
        with TestClient(app, base_url="http://127.0.0.1:8000", headers={"X-KDOG-Request": "1"}) as client:
            checked(client.post("/api/auth/login", json={"username": "pilot", "password": "Synthetic-pilot-only-42"}))
            started = time.perf_counter()
            for index in range(participants):
                item = checked(client.post("/api/cases", json={"event_id": "M6-SYNTHETIC", "participant_id": f"{index + 1:04}",
                    "dog_name": f"가상견-{index + 1:04}"}), 201)
                base = f"/api/cases/{item['case_id']}"
                version = item["manifest"]["sessions"][0]["survey_version"]
                item = checked(client.put(base + "/survey", json={"expected_revision": item["input_revision"],
                    "session_id": item["selected_session_id"], "survey_version": version,
                    "answers": {f"q{question:02}": 3 for question in range(1, 31)}}))
                for camera, video in enumerate(videos):
                    with video.open("rb") as handle:
                        item = checked(client.post(base + "/videos", content=handle, params={
                            "session_id": item["selected_session_id"], "camera_id": f"CAM-{camera + 1}",
                            "filename": "synthetic.mp4", "expected_revision": item["input_revision"]}), 201)
                run = checked(client.post(base + "/analysis", json={"expected_revision": item["input_revision"]}), 202)["runs"][0]
                cases.append((base, run["run_id"]))
            intake_sec = time.perf_counter() - started
            while True:
                run_start = time.perf_counter()
                if not worker.once():
                    break
                durations.append(time.perf_counter() - run_start)
                print(f"synthetic pilot: {len(durations)}/{participants}", flush=True)
            for base, run_id in cases:
                report = checked(client.get(base + "/reports/" + run_id))
                if report["status"] != "ready" or len(report["result"]["scores"]["items"]) != 55:
                    raise RuntimeError("pilot report incomplete")
            export_started = time.perf_counter()
            exported = {}
            for format in ("pdf", "xlsx", "csv"):
                export = checked(client.post("/api/exports", json={"format": format, "event_id": "M6-SYNTHETIC"}), 201)
                if export["count"] != participants:
                    raise RuntimeError("pilot export participant count mismatch")
                checked(client.post(f"/api/exports/{export['export_id']}/generate"))
                download = client.get(f"/api/exports/{export['export_id']}/file")
                if download.status_code != 200 or not download.content:
                    raise RuntimeError("pilot export download failed")
                if format == "xlsx":
                    workbook = load_workbook(BytesIO(download.content), read_only=True)
                    try:
                        ids = {row[1] for row in workbook["전체요약"].iter_rows(min_row=2, values_only=True)}
                        if ids != expected_ids or workbook["행동55항목"].max_row != 55 * participants + 1:
                            raise RuntimeError("pilot workbook omitted participant/items")
                    finally:
                        workbook.close()
                else:
                    with ZipFile(BytesIO(download.content)) as archive:
                        if format == "pdf":
                            if set(archive.namelist()) != {f"M6-SYNTHETIC/{identity}.pdf" for identity in expected_ids}:
                                raise RuntimeError("pilot PDF bundle omitted participant")
                        else:
                            rows = list(csv.reader(StringIO(archive.read("전체요약.csv").decode("utf-8-sig"))))
                            if {row[1] for row in rows[1:]} != expected_ids:
                                raise RuntimeError("pilot CSV bundle omitted participant")
                exported[format] = len(download.content)
            export_sec = time.perf_counter() - export_started
        backup_started = time.perf_counter()
        backup(store, root / "backup", "pilot")
        restore(store, root / "backup", root / "restored")
        restored = summarize(Store(root / "restored"))
        if restored["run_statuses"] != {"scored": participants}:
            raise RuntimeError("pilot restored runs mismatch")
        # Per video: one shared event ledger plus a dog and an owner branch call; one report per run.
        expected_calls = participants * (cameras * 3 + 1)
        usage = summarize(store)
        if usage["calls_reserved"] != expected_calls or len(durations) != participants:
            raise RuntimeError("pilot call count mismatch")
        return {"scenario": "synthetic-only; no external AI calls or actual charges",
            "api_transport": "in-process TestClient; production routes", "media": "640x360 10fps H.264 + AAC color/sine; full decode",
            "participants": participants, "cameras_per_participant": cameras, "seconds_per_video": seconds,
            "total_video_seconds": participants * cameras * seconds, "uploaded_bytes": sum(video.stat().st_size for video in videos) * participants,
            "intake_sec": intake_sec, "worker_total_sec": sum(durations), "worker_median_sec": statistics.median(durations),
            "worker_max_sec": max(durations), "export_sec": export_sec, "export_bytes": exported,
            "backup_restore_sec": time.perf_counter() - backup_started,
            "data_bytes": sum(path.stat().st_size for path in store.root.rglob("*") if path.is_file()),
            "provider_calls_simulated": {"video": len(observer.calls), "evaluate": len(evaluator.calls), "report": len(reporter.calls),
                                         "video_call_shape": "per video: ledger + dog branch + owner branch"},
            "real_provider_calls": 0, "actual_provider_cost": None,
            "platform": platform.platform(), "python": platform.python_version(), "usage": usage}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--participants", type=int, default=20)
    parser.add_argument("--cameras", type=int, default=2)
    parser.add_argument("--seconds", type=int, default=180)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    # Refuse to overwrite earlier measurement evidence.
    with args.output.open("x", encoding="utf-8") as handle:
        json.dump(scenario(args.participants, args.cameras, args.seconds), handle, ensure_ascii=False, indent=2)
