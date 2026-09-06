"""Operator scope, provenance and spreadsheet mapping regression checks."""

from io import BytesIO
import json
import unittest

from openpyxl import Workbook

from app.storage import REPO_ROOT
from tests import test_reporting as reporting


class WorkspaceTests(unittest.TestCase):
    setUp = reporting.ReportingTests.setUp
    login = reporting.ReportingTests.login
    delete_case = reporting.ReportingTests.delete_case
    start = reporting.ReportingTests.start
    view = reporting.ReportingTests.view
    ready_retries = reporting.ReportingTests.ready_retries
    prepare = reporting.ReportingTests.prepare
    report = reporting.ReportingTests.report
    edit = reporting.ReportingTests.edit
    snapshot = reporting.ReportingTests.snapshot
    download = reporting.ReportingTests.download

    def test_preview_selection_freezes_members_and_rejects_changed_results(self):
        self.prepare()
        value = {"format": "xlsx", "case_ids": [self.item["case_id"]]}
        with self.store.connect() as db:
            before = db.execute("SELECT COUNT(*) FROM changes WHERE action='export.snapshot'").fetchone()[0]
        preview = self.client.post('/api/exports/preview', json=value)
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual([m['case_id'] for m in preview.json()['members']], value['case_ids'])
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM changes WHERE action='export.snapshot'").fetchone()[0], before)
        self.assertEqual(self.edit().status_code, 200)
        result = self.client.post('/api/exports', json={**value, 'expected_preview_hash': preview.json()['preview_hash']})
        self.assertEqual(result.status_code, 409, result.text)
        fresh = self.client.post('/api/exports/preview', json=value).json()
        created = self.client.post('/api/exports', json={**value, 'expected_preview_hash': fresh['preview_hash']})
        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(self.client.post('/api/exports', json={**value, 'event_id': 'OTHER'}).status_code, 422)
        self.assertEqual(self.client.post('/api/exports', json={**value, 'case_ids': value['case_ids'] * 2}).status_code, 422)

    def test_delivery_requires_generated_member_and_marks_correction_without_changing_file(self):
        self.prepare()
        export_id = self.snapshot()
        delivery = {'case_id': self.item['case_id'], 'channel': 'email', 'note': '합성 전달 기록'}
        url = f'/api/exports/{export_id}/deliveries'
        self.assertEqual(self.client.post(url, json=delivery).status_code, 409)
        original = self.download(export_id)
        self.assertEqual(self.client.post(url, json={**delivery, 'case_id': 'unrelated'}).status_code, 422)
        saved = self.client.post(url, json=delivery)
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertFalse(saved.json()['members'][0]['correction_needed'])
        self.assertEqual(self.edit().status_code, 200)
        listed = self.client.get('/api/exports').json()
        self.assertTrue(next(e for e in listed if e['export_id'] == export_id)['members'][0]['correction_needed'])
        self.assertEqual(self.download(export_id), original)
        self.login('reviewer')
        self.assertEqual(self.client.post(url, json=delivery).status_code, 403)
        self.login('operator')
        self.delete_case()
        self.assertEqual(self.client.post(url, json=delivery).status_code, 403)

    def test_custom_headers_and_horizontal_original_require_explicit_connections(self):
        mapped = {'columns': {'event_id': '행사', 'participant_id': '번호', 'dog_name': '이름'}}
        response = self.client.post('/api/imports/preview', params={'kind': 'participants', 'format': 'csv', 'mapping': json.dumps(mapped)},
            content='행사,번호,이름\nMAPPING,0007,가상견\n'.encode('utf-8'))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['rows'][0]['participant_id'], '0007')
        self.assertEqual(self.client.post('/api/imports/preview', params={'kind': 'participants', 'format': 'csv', 'mapping': '{'}, content=b'a').status_code, 422)
        book = Workbook()
        sheet = book.active
        sheet.title = '데이터입력'
        catalog = json.loads((REPO_ROOT / 'resources/catalogs/survey-v1.json').read_text(encoding='utf-8'))
        sheet.cell(5, 6, 'ID로 자동 연결하지 않는 이름')
        for index, item in enumerate(catalog['items'], 6):
            sheet.cell(index, 1, index - 5)
            sheet.cell(index, 4, item['text'])
            sheet.cell(index, 6, 3)
        output = BytesIO()
        book.save(output)
        columns = self.client.post('/api/imports/columns', params={'format': 'xlsx', 'sheet': '데이터입력', 'horizontal': True}, content=output.getvalue())
        self.assertEqual(columns.status_code, 200, columns.text)
        self.assertEqual(columns.json()['columns'][0]['key'], 'F')
        mapping = {'sheet': '데이터입력', 'horizontal': {'F': self.item['case_id']}}
        preview = self.client.post('/api/imports/preview', params={'kind': 'survey', 'format': 'xlsx', 'mapping': json.dumps(mapping)}, content=output.getvalue())
        self.assertEqual(preview.status_code, 200, preview.text)
        row = preview.json()['rows'][0]
        self.assertEqual(row['case_id'], self.item['case_id'])
        self.assertEqual(row['session_label'], '1차 촬영')
        self.assertEqual(row['source_location'], '데이터입력 F6:F35')
        self.assertEqual(len(row['survey']['answers']), 30)
        sheet.cell(6, 6, 9)
        invalid = BytesIO()
        book.save(invalid)
        errors = self.client.post('/api/imports/preview', params={'kind': 'survey', 'format': 'xlsx', 'mapping': json.dumps(mapping)}, content=invalid.getvalue()).json()['errors']
        self.assertIn('데이터입력 F6:F35: q01', errors[0])
        sheet.cell(6, 4, '다른 질문')
        changed = BytesIO()
        book.save(changed)
        self.assertEqual(self.client.post('/api/imports/preview', params={'kind': 'survey', 'format': 'xlsx', 'mapping': json.dumps(mapping)}, content=changed.getvalue()).status_code, 422)
        book.close()

    def test_identity_and_checklist_are_versioned_and_reviewer_cannot_change_them(self):
        self.prepare()
        export_id = self.snapshot()
        old = self.download(export_id)
        item = self.client.get(self.base).json()
        value = {'expected_revision': item['input_revision'], 'participant_id': '0000008', 'dog_name': '정정 가상견', 'reservation_at': ''}
        changed = self.client.put(self.base, json=value)
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()['participant_id'], '0000008')
        self.assertIsNone(changed.json()['manifest']['display_run_id'])
        self.assertEqual(self.download(export_id), old)
        self.assertEqual(self.client.put(self.base, json=value).status_code, 409)
        metadata = {'expected_revision': changed.json()['input_revision'], 'capture_mode': 'sequential', 'route_note': '합성 촬영', 'checklist': {'entry': 'performed', 'training': 'skipped'}}
        session_url = f"{self.base}/sessions/{item['selected_session_id']}"
        saved = self.client.put(session_url, json=metadata)
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(saved.json()['manifest']['sessions'][0]['checklist']['training'], 'skipped')
        self.login('reviewer')
        self.assertEqual(self.client.put(self.base, json=value).status_code, 403)
        self.assertEqual(self.client.put(session_url, json=metadata).status_code, 403)

    def test_narration_round_trip_keeps_other_paragraphs_and_before_after_history(self):
        view = self.prepare()
        original = view['report']
        narration = {'cover': {'text': '표지 문단만 정정합니다.', 'evidence_ids': []},
            'comments': [{'text': d['comment'], 'evidence_ids': d['evidence_ids']} for d in original['domains']],
            'cross': {'text': original['cross_type']['explanation'], 'evidence_ids': original['cross_type']['evidence_ids']},
            'tips': original['tips']}
        response = self.client.put(self.report_base, json={'expected_revision': view['revision'], 'expected_source_hash': view['source_hash'],
            'reason': '표지 수정', 'narration': narration})
        self.assertEqual(response.status_code, 200, response.text)
        saved = response.json()
        for key in ('domains', 'cross_type', 'tips'):
            self.assertEqual(saved['report'][key], original[key])
        self.assertEqual(saved['history'][-1]['before'], original)
        self.assertEqual(saved['history'][-1]['after'], saved['report'])
