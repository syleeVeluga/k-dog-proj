"""05 inputs against hand-calculated 03 formula values plus explicit R4–R6 interpretation states.

This is not an Excel-engine recalculation test. R5 makes pair 2 invalid despite a blank phase count;
03 여러쌍비교!Z6 checks the blank first. Customer confirmation remains pending.
"""

import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.domain.catalog import BehaviorCatalog
from app.domain.contracts import Rater, ScoreSheet
from app.import_catalogs import CUSTOMER_DIR
from app.scoring import behavior_scores

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = CUSTOMER_DIR / "05_예비촬영_6쌍_채점_20260913.xlsx"
PAIRS = ("천우미", "김지유", "배보경", "최선미", "이하연", "송소연")
NONE = {"mean": None, "lean": None, "width": None, "degree": None}
EMPTY_DOMAINS = {domain: NONE for domain in ("EDU", "SOC_E", "SOC_H", "ATT", "SYN", "EXIT")}
MISSING = ("missing", None)
# Numeric expectations follow 03 여러쌍비교: D/G/P/S lean = ROUND(mean(BI)-3, 2); E/H/K/N/Q/T width = max|x-3| over
# BI and ONE items; F/I/L/O/R/U degree = ROUND(mean(ONE), 2); J/M types; V/W/X distance differences; Z = phases/6.
# Inputs: 05 sheets 1–3, J:O and each catalog item's source_row. R4 missing-material types and R5 invalid/excluded
# states are document/label interpretations, not literal formula results. R6 does not fire (05 J12:O12 blank).
EXPECTED = {
    "천우미": {
        "domains": {
            "EDU": {"mean": 3.0, "lean": 0.0, "width": 0, "degree": 3.0},      # OWN-02 = 3 (BI), OWN-04 = 3 (ONE)
            "SOC_E": {"mean": 2.8, "lean": -0.2, "width": 1, "degree": None},  # 3,3,2 + 3,3
            "SOC_H": {"mean": 2.75, "lean": -0.25, "width": 1, "degree": None},  # 2 + 2,3,4
            "ATT": {"mean": 2.43, "lean": -0.57, "width": 1, "degree": None},  # 2 + 2,2,2,3,3,3 = 17/7 (무시 2항목 빈칸)
            "SYN": {"mean": 2.67, "lean": -0.33, "width": 1, "degree": None},  # 3 + 2,3
            "EXIT": {"mean": 3.0, "lean": 0.0, "width": 0, "degree": None},    # 3 + 3
        },
        "indicators": {"adaptation": MISSING, "recovery": ("calculated", 1.0), "stranger_calming": ("calculated", 1.0),
                       "sync_rate": ("calculated", 5 / 6)},
        "types": {"attachment": MISSING, "sociability_person": ("calculated", "담담·거리둠")},
        "baseline": None,
    },
    "김지유": {  # 05 3_보호자행동!K13=3, 2_개행동!K25 blank: R5 invalid-first; literal 03 Z6 would be blank.
        "domains": {**EMPTY_DOMAINS, "SOC_E": {"mean": 3.0, "lean": 0.0, "width": 0, "degree": None},
                    "ATT": {"mean": 3.0, "lean": 0.0, "width": 0, "degree": None}},
        "indicators": {"adaptation": MISSING, "recovery": MISSING, "stranger_calming": MISSING, "sync_rate": ("invalid", None)},
        "types": {"attachment": MISSING, "sociability_person": MISSING},
        "baseline": None,
    },
    "배보경": {  # DOG-01 = 3, DOG-05 = 4, DOG-06 = 3
        "domains": {**EMPTY_DOMAINS, "SOC_E": {"mean": 3.0, "lean": 0.0, "width": 0, "degree": None},
                    "ATT": {"mean": 3.5, "lean": 0.5, "width": 1, "degree": None}},
        "indicators": {key: MISSING for key in ("adaptation", "recovery", "stranger_calming", "sync_rate")},
        "types": {"attachment": MISSING, "sociability_person": MISSING},
        "baseline": None,
    },
    "최선미": {
        "domains": EMPTY_DOMAINS,
        "indicators": {key: MISSING for key in ("adaptation", "recovery", "stranger_calming", "sync_rate")},
        "types": {"attachment": MISSING, "sociability_person": MISSING},
        "baseline": None,
    },
}
EXPECTED["이하연"] = EXPECTED["송소연"] = {  # DOG-01 = 3, OWN-06 = 3
    **EXPECTED["최선미"], "domains": {**EMPTY_DOMAINS, "SOC_E": {"mean": 3.0, "lean": 0.0, "width": 0, "degree": None}},
}


def read_pairs(catalog):
    """Column J..O of the three item sheets, mapped to catalog ids by row (row = number + 4); blanks are unreadable."""
    book = load_workbook(WORKBOOK, data_only=True)
    try:
        sheets = {}
        for index, name in enumerate(PAIRS):
            column = "JKLMNO"[index]
            entries = []
            for item in catalog.rated_items():
                cell = book[item.sheet][f"{column}{item.source_row}"]
                if cell.value is None:
                    entries.append({"item_id": item.item_id, "score": None, "status": "unreadable", "reason": "05 빈칸"})
                else:
                    entries.append({"item_id": item.item_id, "score": int(cell.value), "status": "scored", "reason": None})
            header = book["2_개행동"][f"{column}4"].value
            if not str(header).endswith(name):
                raise ValueError(f"pair column {column} is {header!r}, expected {name}")
            sheets[name] = ScoreSheet(sheet_id=f"pair-{index + 1}", case_id=f"case-{index + 1}", session_id=f"session-{index + 1}",
                                      catalog_version=catalog.version, rater=Rater(rater_id="human-05", kind="human"),
                                      recorded_at="2026-09-13T20:00:00+09:00", items=tuple(entries))
        return sheets
    finally:
        book.close()


class GoldenTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = BehaviorCatalog.model_validate_json((ROOT / "resources/catalogs/behavior-v2.json").read_bytes())
        cls.sheets = read_pairs(cls.catalog)

    def test_source_cells_distinguish_formula_from_interpretation(self):
        book = load_workbook(WORKBOOK, data_only=False)
        try:
            self.assertEqual(book["3_보호자행동"]["K13"].value, 3)
            self.assertIsNone(book["2_개행동"]["K25"].value)
            self.assertEqual([book["2_개행동"][f"{c}20"].value for c in "JKLMNO"], [0, None, None, None, None, None])
            self.assertEqual([book["3_보호자행동"][f"{c}12"].value for c in "JKLMNO"], [None] * 6)
        finally:
            book.close()
        book = load_workbook(CUSTOMER_DIR / "03_행동_채점표_42항목_20260913.xlsx", data_only=False)
        try:
            self.assertEqual(book["여러쌍비교"]["Z6"].value,
                             '=IF(\'2_개행동\'!$J$25="","",IF(\'3_보호자행동\'!$J$13=3,"무효",\'2_개행동\'!$J$25/6))')
            self.assertEqual([book["3_보호자행동"][f"{c}4"].value for c in "CDE"], ["1", "2", "3"])
            self.assertIn("무시 항목 무효", book["3_보호자행동"]["E12"].value)
            self.assertIn("걷기 항목과 동조율 무효", book["3_보호자행동"]["E13"].value)
        finally:
            book.close()

    def test_pair_inputs_are_the_filled_cells_of_the_workbook(self):
        # 00 문서의 「35 / 42」는 참고 열 표기이고, 05 시트 1~3의 채점 대상 41칸 중 채워진 칸은 아래와 같다.
        scored = {name: sum(item.status == "scored" for item in sheet.items) for name, sheet in self.sheets.items()}
        self.assertEqual(scored, {"천우미": 33, "김지유": 5, "배보경": 4, "최선미": 0, "이하연": 2, "송소연": 2})

    def test_every_pair_matches_the_workbook_formulas(self):
        for name, sheet in self.sheets.items():
            result = behavior_scores(sheet, self.catalog)
            expected = EXPECTED[name]
            with self.subTest(pair=name, part="domains"):
                self.assertEqual({d.domain: {"mean": d.mean, "lean": d.lean, "width": d.width, "degree": d.degree} for d in result.domains},
                                 expected["domains"])
            with self.subTest(pair=name, part="indicators"):
                self.assertEqual({i.key: (i.status, i.value) for i in result.indicators}, expected["indicators"])
            with self.subTest(pair=name, part="types"):
                self.assertEqual({t.key: (t.status, t.label) for t in result.types}, expected["types"])
            with self.subTest(pair=name, part="baseline"):
                self.assertEqual(result.baseline_arousal, expected["baseline"])

    def test_walk_invalidation_is_reported_not_hidden(self):
        result = behavior_scores(self.sheets["김지유"], self.catalog)
        syn = next(d for d in result.domains if d.domain == "SYN")
        self.assertEqual((syn.excluded_count, syn.scored_count, syn.unreadable_count), (3, 0, 0))
        sync = next(i for i in result.indicators if i.key == "sync_rate")
        self.assertIn("OWN-09 = 3", sync.reason)
        self.assertEqual(len(result.items), 41)
        self.assertEqual(result.scoring_rule_version, "scoring-v2")


if __name__ == "__main__":
    unittest.main()
