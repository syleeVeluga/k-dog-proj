"""M4 revision, grounded narration, export isolation and interruption regression tests."""

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path
import subprocess
import threading
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

from app import exports
from app.analysis import write_output
from app.gemini import ProviderError
from app.report_models import ExportRequest, Narration
from app.reporting import Reporter, active_report_configuration
from app.storage import encode
from tests import test_observation as observation


class ReportingTests(unittest.TestCase):
    setUp = observation.ObservationTests.setUp
    login = observation.ObservationTests.login
    access = observation.ObservationTests.access
    start = observation.ObservationTests.start
    view = observation.ObservationTests.view
    ready_retries = observation.ObservationTests.ready_retries

    def prepare(self):
        self.run_id = self.start()
        self.report_base = self.base + "/reports/" + self.run_id
        self.worker.once()
        return self.report()

    def report(self):
        response = self.client.get(self.report_base)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def edit(self, *, option=None, revision=None):
        view = self.report()
        item = next(i for a in view["result"]["evaluations"] for i in a["evaluation"]["items"] if i["item_id"] == "BS-01")
        selected = option or next(c for c in view["result"]["behavior_items"] if c["item_id"] == "BS-01")["options"][-1]["option_id"]
        return self.client.put(self.report_base, json={"expected_revision": revision or view["revision"], "expected_source_hash": view["source_hash"], "reason": "원본 근거 재검토",
            "item": {**item, "selected_option_id": selected, "status": "scored", "reason": "원본 근거 재검토"}})

    def snapshot(self, format="xlsx", **kwargs):
        value = {"format": format, "case_id": self.item["case_id"], "run_id": self.run_id, **kwargs}
        response = self.client.post("/api/exports", json=value)
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["export_id"]

    def download(self, export_id):
        response = self.client.post(f"/api/exports/{export_id}/generate")
        self.assertEqual(response.status_code, 200, response.text)
        result = self.client.get(f"/api/exports/{export_id}/file")
        self.assertEqual(result.status_code, 200, result.text[:200] if result.status_code != 200 else "")
        return result.content

    def test_report_does_not_block_scores_and_binds_references_without_invented_mapping(self):
        entered, release = threading.Event(), threading.Event()
        reporter = self.worker.reporter
        def delayed(config, context, guard):
            entered.set()
            self.assertTrue(release.wait(8))
            reporter.callback = None
            return reporter.write(config, context, guard)
        reporter.callback = delayed
        self.run_id = self.start()
        self.report_base = self.base + "/reports/" + self.run_id
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(self.worker.once)
            try:
                self.assertTrue(entered.wait(8))
                view = self.report()
                self.assertEqual(len(view["result"]["scores"]["items"]), 55)
                self.assertEqual(view["status"], "running")
                self.assertFalse(future.done())
            finally:
                release.set()
            future.result(timeout=10)
        view = self.report()
        self.assertEqual(view["status"], "ready")
        self.assertEqual([d["value"] for d in view["report"]["domains"]], [None] * 4)
        self.assertIsNone(view["report"]["cross_type"]["type_name"])
        context = reporter.calls[0][1]
        for key in ("dog_name", "participant_id", "case_id", "run_id", "storage_ref"):
            self.assertNotIn('"' + key + '"', encode(context))

    def test_score_edit_invalidates_text_preserves_ai_and_rejects_stale_or_foreign_options(self):
        initial = self.prepare()
        original = self.view()["scores"]
        self.login("reviewer")
        updated = self.edit()
        self.assertEqual(updated.status_code, 200, updated.text)
        view = updated.json()
        self.assertEqual(view["status"], "stale")
        self.assertIsNone(view["report"])
        self.assertEqual(view["revision"], 2)
        self.assertEqual(self.view()["scores"], original)
        self.assertNotEqual(view["result"]["scores"], original)
        self.assertEqual(self.edit(revision=1).status_code, 409)
        self.assertEqual(self.edit(option="invented").status_code, 422)
        self.assertEqual(self.report()["revision"], 2)
        self.assertEqual(initial["revision"], 1)
        count = len(self.worker.evaluator.calls)
        self.assertEqual(self.client.post(self.report_base + "/generate").status_code, 200)
        self.worker.once()
        self.assertEqual(len(self.worker.evaluator.calls), count)
        self.assertEqual(self.report()["status"], "ready")
        self.assertEqual(self.report()["report"]["result_revision"], 2)

    def test_null_edit_and_manual_explanation_require_reason_and_valid_evidence(self):
        self.prepare()
        view = self.report()
        item = next(i for a in view["result"]["evaluations"] for i in a["evaluation"]["items"] if i["item_id"] == "BS-01")
        response = self.client.put(self.report_base, json={"expected_revision": 1, "expected_source_hash": view["source_hash"], "reason": "가림 재검토", "item": {
            **item, "status": "not_visible", "selected_option_id": None, "reason": "가림"}})
        self.assertEqual(response.status_code, 200, response.text)
        score = response.json()["result"]["scores"]["items"][0]
        self.assertIsNone(score["raw_score"])
        part = {"text": "관찰 조건을 확인하고 짧게 연습해 보세요.", "evidence_ids": ["other-case-evidence"]}
        narration = {"cover": part, "comments": [part] * 4, "cross": part, "tips": [part]}
        body = {"expected_revision": 2, "expected_source_hash": response.json()["source_hash"], "reason": "설명 정정", "narration": narration}
        self.assertEqual(self.client.put(self.report_base, json=body).status_code, 422)
        part["evidence_ids"] = [view["result"]["evidence"][0]["evidence_id"]]
        self.assertEqual(self.client.put(self.report_base, json={**body, "reason": " "}).status_code, 422)
        response = self.client.put(self.report_base, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["status"], "manual")
        self.assertEqual(response.json()["history"][-1]["reason"], "설명 정정")

    def test_export_snapshot_survives_edits_and_redownload_never_calls_ai(self):
        self.prepare()
        before = self.snapshot()
        self.assertEqual(self.edit().status_code, 200)
        after = self.snapshot()
        first = self.download(before)
        second = self.download(after)
        one = load_workbook(BytesIO(first))
        two = load_workbook(BytesIO(second))
        self.assertNotEqual(one["행동55항목"]["H2"].value, two["행동55항목"]["H2"].value)
        self.assertEqual(one["전체요약"]["F2"].value, 1)
        self.assertEqual(two["전체요약"]["F2"].value, 2)
        self.assertEqual(two["전체요약"]["H2"].value, "stale")
        self.assertNotIn("가상 시험: 입장", str(list(two["리포트"].values)))
        count = len(self.worker.reporter.calls)
        self.assertEqual(self.download(before), first)
        self.assertEqual(len(self.worker.reporter.calls), count)
        self.assertEqual(one["전체요약"]["B2"].value, "0001")
        self.assertEqual(one["전체요약"]["B2"].data_type, "s")
        self.assertEqual(one["행동55항목"].max_row, 56)
        self.assertEqual(one["설문"].max_row, 31)

    def test_all_exports_include_unprocessed_participant_without_cross_case_leak(self):
        self.prepare()
        other = self.client.post("/api/cases", json={"event_id": "M2", "participant_id": "0002", "dog_name": "타참가자고유이름"}).json()
        consent = self.item["consent"]
        self.assertEqual(self.client.put(f"/api/cases/{other['case_id']}/access", json={"expected_revision": 1, "consent": consent, "deletion_requested": False}).status_code, 200)
        individual = self.download(self.snapshot())
        wb = load_workbook(BytesIO(individual))
        self.assertNotIn("타참가자고유이름", str([list(ws.values) for ws in wb]))
        all_id = self.snapshot(case_id=None, run_id=None, event_id="M2")
        wb = load_workbook(BytesIO(self.download(all_id)))
        self.assertEqual(wb["전체요약"].max_row, 3)
        self.assertEqual(wb["행동55항목"].max_row, 111)
        self.assertEqual(wb["설문"].max_row, 61)
        self.assertEqual(wb["전체요약"]["D3"].value, "not_started")
        archive = ZipFile(BytesIO(self.download(self.snapshot("pdf", case_id=None, run_id=None, event_id="M2"))))
        self.assertEqual(set(archive.namelist()), {"M2/0001.pdf", "M2/0002.pdf"})
        first = "".join(p.extract_text() for p in PdfReader(BytesIO(archive.read("M2/0001.pdf"))).pages)
        self.assertIn("가상견", first)
        self.assertNotIn("타참가자고유이름", first)

    def test_pdf_korean_csv_zip_and_excel_literal_formula_text(self):
        self.prepare()
        snapshot = exports.capture(self.store, ExportRequest(format="xlsx", case_id=self.item["case_id"], run_id=self.run_id), "operator")
        snapshot["members"][0]["dog_name"] = '=HYPERLINK("https://example.invalid")'
        workbook = load_workbook(BytesIO(exports.workbook(snapshot)))
        self.assertEqual(workbook["전체요약"]["C2"].data_type, "s")
        pdf = PdfReader(BytesIO(self.download(self.snapshot("pdf"))))
        text = "\n".join(page.extract_text() for page in pdf.pages)
        for value in ("가상견", "4영역", "오늘의 팁", "진단이 아닌", "0001", "②④"):
            self.assertIn(value, text)
        snapshot["format"] = "csv"
        archive = ZipFile(BytesIO(exports.render(snapshot)[0]))
        self.assertEqual(len(archive.namelist()), 9)
        self.assertIn("'=HYPERLINK", archive.read("전체요약.csv").decode("utf-8-sig"))

    def test_permission_and_consent_checks_cover_snapshot_generation_and_download(self):
        self.prepare()
        export_id = self.snapshot()
        self.download(export_id)
        self.login("developer")
        for path in (self.report_base, self.report_base + "/image", "/api/exports", f"/api/exports/{export_id}/file"):
            self.assertEqual(self.client.get(path).status_code, 403)
        self.assertEqual(self.client.post(f"/api/exports/{export_id}/generate").status_code, 403)
        self.login("operator")
        self.access(False)
        self.assertEqual(self.client.get(self.report_base).status_code, 403)
        self.assertEqual(self.client.post(f"/api/exports/{export_id}/generate").status_code, 403)
        self.assertEqual(self.client.get(f"/api/exports/{export_id}/file").status_code, 403)
        self.assertEqual(self.client.get("/api/exports").json(), [])

    def test_withdrawal_during_render_does_not_adopt_file(self):
        self.prepare()
        export_id = self.snapshot()
        original = exports.render
        def revoked(snapshot):
            raw = original(snapshot)
            self.access(False)
            return raw
        with patch("app.exports.render", side_effect=revoked):
            self.assertEqual(self.client.post(f"/api/exports/{export_id}/generate").status_code, 403)
        with self.store.connect() as db:
            self.assertIsNone(db.execute("SELECT 1 FROM changes WHERE target=? AND action='export.file'", (export_id,)).fetchone())

    def test_interrupted_export_publication_recovers_frozen_snapshot(self):
        self.prepare()
        export_id = self.snapshot()
        original = self.store.audit
        def interrupted(db, actor, target, action, detail):
            if action == "export.file":
                raise OSError("synthetic power interruption")
            return original(db, actor, target, action, detail)
        with patch.object(self.store, "audit", side_effect=interrupted):
            self.assertEqual(self.client.post(f"/api/exports/{export_id}/generate").status_code, 503)
        self.assertEqual(self.edit().status_code, 200)
        wb = load_workbook(BytesIO(self.download(export_id)))
        self.assertEqual(wb["전체요약"]["F2"].value, 1)

    def test_report_failure_retry_and_revoked_late_response_preserve_scores(self):
        self.worker.reporter.callback = lambda *a: (_ for _ in ()).throw(ProviderError("developer_settings_required"))
        self.prepare()
        self.assertEqual(self.view()["status"], "scored")
        self.assertEqual(self.report()["status"], "failed")
        self.worker.reporter.callback = None
        self.client.post(self.report_base + "/generate")
        self.worker.once()
        self.assertEqual(self.report()["status"], "ready")
        self.edit()
        reporter = self.worker.reporter
        def revoked(config, context, guard):
            reporter.callback = None
            response = reporter.write(config, context, guard)
            self.access(False)
            return response
        reporter.callback = revoked
        self.client.post(self.report_base + "/generate")
        self.worker.once()
        self.assertEqual(self.view()["scores"], None)
        self.assertEqual(self.view()["status"], "stopped")

    def test_report_schema_repair_limit_and_older_retry_does_not_loop_new_revision(self):
        reporter = self.worker.reporter
        def invalid(config, context, guard):
            raise ProviderError("evaluation_schema_invalid", retryable=True)
        reporter.callback = invalid
        self.prepare()
        self.ready_retries()
        self.worker.once()
        self.assertEqual(len(reporter.calls), 2)
        self.assertEqual(self.report()["status"], "failed")
        self.assertEqual(self.client.post(self.report_base + "/generate").status_code, 409)
        self.assertEqual(self.edit().status_code, 200)
        reporter.callback = None
        self.client.post(self.report_base + "/generate")
        self.worker.once()
        self.assertEqual(self.report()["status"], "ready")
        self.assertEqual(self.view()["status"], "scored")

    def test_report_file_published_before_db_adoption_recovers_without_recall(self):
        original = write_output
        self.run_id = self.start()
        self.report_base = self.base + "/reports/" + self.run_id
        def interrupted(store, row, step, payload):
            result = original(store, row, step, payload)
            if "report" in payload:
                raise OSError("synthetic interruption")
            return result
        with patch("app.worker.write_output", side_effect=interrupted):
            self.worker.once()
        count = len(self.worker.reporter.calls)
        self.client.post(self.report_base + "/generate")
        self.worker.once()
        self.assertEqual(self.report()["status"], "ready")
        self.assertEqual(len(self.worker.reporter.calls), count)

    def test_original_frame_membership_time_hash_and_snapshot_image(self):
        self.prepare()
        video = self.item["manifest"]["sessions"][0]["videos"][0]
        # Re-register a real generated clip; original fake bytes must never render a frame.
        clip = Path(self.temp.name) / "generated.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-nostdin", "-f", "lavfi", "-i", "color=c=teal:s=320x240:d=4", "-c:v", "libx264", str(clip)], check=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        item = self.client.post(self.base + "/videos", content=clip.read_bytes(), params={"session_id": self.item["selected_session_id"], "camera_id": "CAM-3", "filename": "generated.mp4", "expected_revision": self.item["input_revision"]}).json()
        self.item = item
        self.prepare()
        video = item["manifest"]["sessions"][0]["videos"][-1]
        payload = {"expected_revision": 1, "video_id": video["video_id"], "second": 1.0}
        for bad in ({"video_id": "foreign"}, {"second": 4.0}, {"second": -1.0}):
            self.assertEqual(self.client.put(self.report_base + "/image", json={**payload, **bad}).status_code, 422)
        response = self.client.put(self.report_base + "/image", json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["status"], "ready")
        self.assertTrue(self.client.get(self.report_base + "/image").content.startswith(b"\x89PNG"))
        wb = load_workbook(BytesIO(self.download(self.snapshot())))
        self.assertEqual(len(wb["대표이미지1"]._images), 1)
        exported_pdf = PdfReader(BytesIO(self.download(self.snapshot("pdf"))))
        self.assertGreater(sum(len(page.images) for page in exported_pdf.pages), 0)

    def test_config_permission_version_and_existing_run_pinning(self):
        self.prepare()
        self.assertEqual(self.client.get("/api/developer/report").status_code, 403)
        self.login("developer")
        settings = self.client.get("/api/developer/report").json()
        body = {"expected_version": settings["version"], "selection": {"provider": "anthropic", "model": "test-only"}}
        self.assertEqual(self.client.put("/api/developer/report", json=body).status_code, 200)
        self.assertEqual(self.client.put("/api/developer/report", json=body).status_code, 409)
        with self.store.connect() as db:
            old = json.loads(db.execute("SELECT config_snapshot_json FROM runs WHERE run_id=?", (self.run_id,)).fetchone()[0])
            self.assertEqual(old["report"]["provider"], "gemini")
            self.assertEqual(active_report_configuration(db, "")[1]["provider"], "anthropic")

    def test_cold_review_late_branch_cannot_adopt_text_drafted_on_partial_scores(self):
        from tests.evaluation_fixtures import evaluation_response
        def partial(config, context, guard):
            if context["branch"] == "owner":
                raise ProviderError("provider_http_400")
            return evaluation_response(context), {}
        self.worker.evaluator.callback = partial
        view = self.prepare()
        part = {"text": "부분 결과만 참고한 수동 초안입니다.", "evidence_ids": []}
        body = {"expected_revision": view["revision"], "expected_source_hash": view["source_hash"], "reason": "부분 결과 검토",
                "narration": {"cover": part, "comments": [part] * 4, "cross": part, "tips": [part]}}
        self.worker.evaluator.callback = None
        self.client.post(self.base + f"/analysis/{self.run_id}/retry")
        self.worker.once()
        self.assertEqual(self.report()["revision"], view["revision"])
        self.assertEqual(self.client.put(self.report_base, json=body).status_code, 409)

    def test_cold_review_pdf_bundle_names_cannot_collide_between_events(self):
        self.prepare()
        snapshot = exports.capture(self.store, ExportRequest(format="pdf", case_id=self.item["case_id"], run_id=self.run_id), "operator")
        one = snapshot["members"][0]
        one.update(event_id="A-B", participant_id="C")
        snapshot["members"].append({**one, "event_id": "A", "participant_id": "B-C"})
        snapshot["individual"] = False
        archive = ZipFile(BytesIO(exports.render(snapshot)[0]))
        self.assertEqual(len(set(archive.namelist())), 2)

    def test_cold_review_queued_run_exports_original_survey_before_worker(self):
        session = self.item["manifest"]["sessions"][0]
        response = self.client.put(self.base + "/survey", json={"expected_revision": self.item["input_revision"],
            "session_id": session["session_id"], "survey_version": session["survey_version"], "answers": {f"q{i:02}": 3 for i in range(1, 31)}})
        self.assertEqual(response.status_code, 200, response.text)
        self.item = response.json()
        self.run_id = self.start()
        wb = load_workbook(BytesIO(self.download(self.snapshot())))
        self.assertEqual([wb["설문"].cell(row, 5).value for row in range(2, 32)], [3] * 30)


class ReportAdapterTests(unittest.TestCase):
    def test_three_provider_structured_narration_contracts(self):
        from tests.report_fixtures import FakeReporter
        response, _ = FakeReporter().write({}, {"evidence": []}, lambda: None)
        text = response.model_dump_json()
        bodies = {
            "gemini": {"status": "completed", "steps": [{"type": "model_output", "content": [
                {"type": "text", "text": text}]}]},
            "openai": {"status": "completed", "output": [{"type": "message", "content": [{"type": "output_text", "text": text}]}]},
            "anthropic": {"stop_reason": "end_turn", "content": [{"type": "text", "text": text}]},
        }
        with patch.dict("os.environ", {"GEMINI_API_KEY": "fake", "OPENAI_API_KEY": "fake", "ANTHROPIC_API_KEY": "fake"}):
            for provider, body in bodies.items():
                with self.subTest(provider=provider), patch("app.evaluation.request", return_value=(body, {})) as request:
                    result, _ = Reporter().write({"provider": provider, "model": "synthetic", "prompt": "test", "max_output_tokens": 8192}, {"scores": {}}, lambda: None)
                    self.assertIsInstance(result, Narration)
                    encoded = encode(request.call_args.kwargs["data"])
                    self.assertIn('"comments"', encoded)
                    self.assertNotIn('"branch"', encoded)


if __name__ == "__main__":
    unittest.main()
