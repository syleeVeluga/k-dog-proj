"""Read exported files back using synthetic pinned S1 observations only."""
import csv
from io import BytesIO, StringIO
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook

from app import exports_v4 as exports, sheets_v4 as sheets, maintenance
from app.domain.exports_v4 import ExportSnapshotV4
from app.domain.sheets_v4 import SheetReferenceV4, SheetDocumentV4
from app.storage import Store, encode, uid
from app.domain.sheets_v4 import AI_ACCOUNT
from tests.test_opinions_v4 import model
from tests import test_report_runs_v4 as report_fixture


class ExportV4Tests(unittest.TestCase):
    def setUp(self):
        self.fixture = report_fixture.ReportRunV4Tests(); self.fixture.setUp(); self.addCleanup(self.fixture.doCleanups)
        self.store, self.user = self.fixture.store, self.fixture.user
        self.case_id, self.session_id = self.fixture.case_id, self.fixture.session_id
        with self.store.connect(write=True) as db:
            db.execute("UPDATE cases SET consent_confirmed=1 WHERE case_id=?", (self.case_id,))
        self.request = model(exports.ExportCreateV4, request_id="synthetic-export", format="csv_zip", reason="synthetic research selection",
            members=[{"case_id":self.case_id,"session_id":self.session_id,
                "sheet":self.fixture.basic["document"]["input"],"basic":self.fixture.basic_ref,"final":self.fixture.final["reference"]}])

    def create(self, request=None):
        return exports.create(self.store, request or self.request, self.user)

    def csv_tables(self, data):
        with ZipFile(BytesIO(data)) as archive:
            return {name[:-4]:list(csv.DictReader(StringIO(archive.read(name).decode("utf-8-sig")))) for name in archive.namelist()}

    def snapshot(self, export_id):
        with self.store.connect() as db:return exports._load(self.store,exports._record(db,export_id))

    def test_csv_is_long_s1_only_with_exact_types_nulls_and_separate_results(self):
        created = self.create()
        data,mime,_ = exports.download(self.store,created["export_id"],self.user)
        tables = self.csv_tables(data)
        raw = tables["raw_observations"]
        self.assertEqual(len(raw),90)
        self.assertEqual(sum(row["input_kind"] == "numeric" for row in raw),83)
        self.assertEqual(sum(row["input_kind"] == "automatic" for row in raw),4)
        self.assertEqual(sum(row["input_kind"] == "memo" for row in raw),3)
        self.assertNotIn("개32",{row["code"] for row in raw})
        self.assertEqual(next(row for row in raw if row["code"] == "보23")["value"],"2")
        self.assertEqual(next(row for row in raw if row["code"] == "개5")["value"],"")
        self.assertTrue(tables["automatic_calculations"])
        self.assertTrue(tables["final_results"])
        self.assertEqual(len(tables["survey_responses"]),28)
        self.assertEqual(mime,"application/zip")
        for forbidden in ("guardian_name","dog_name","participant_id","assigned_username","api_key"):
            self.assertNotIn(forbidden,raw[0])

    def test_xlsx_keeps_numeric_zero_negative_and_sanitizes_formula_text(self):
        snapshot = exports._snapshot(self.store,self.request,self.user)
        base = exports.tables(snapshot)
        base["raw_observations"][0]["value"] = 0
        base["raw_observations"][1]["value"] = -2
        base["raw_observations"][2]["review_memo"] = "\t=HYPERLINK(\"https://invalid\")"
        snapshot = snapshot.model_copy(update={"format":"xlsx"})
        with patch.object(exports,"tables",return_value=base):data = exports.render(snapshot)
        workbook = load_workbook(BytesIO(data),read_only=True,data_only=False)
        rows = list(workbook["raw_observations"].values); index = rows[0].index("value"); note = rows[0].index("review_memo")
        self.assertEqual(rows[1][index],0); self.assertEqual(rows[2][index],-2)
        self.assertTrue(rows[3][note].startswith("'\t="))
        self.assertEqual(workbook["raw_observations"].cell(3,index+1).data_type,"n")
        workbook.close()
        for note in ("x"*32768,"x\x00y"):
            base["raw_observations"][0]["review_memo"] = note
            with patch.object(exports,"tables",return_value=base),self.assertRaises(HTTPException):exports.render(snapshot)
            with patch.object(exports,"tables",return_value=base):preserved = exports.render(snapshot.model_copy(update={"format":"csv_zip"}))
            self.assertEqual(self.csv_tables(preserved)["raw_observations"][0]["review_memo"],note)

    def test_fixed_snapshot_idempotency_and_later_input_edit_preserve_download(self):
        created = self.create(); self.assertEqual(self.create()["export_id"],created["export_id"])
        before,_,_ = exports.download(self.store,created["export_id"],self.user)
        with self.store.connect(write=True) as db:
            case = self.store.case(db,self.case_id); manifest = self.store.manifest(case)
            manifest.sessions[0].survey["s01"] = 3
            self.store.save(db,case,manifest,self.user.username,"synthetic.survey.update")
        after,_,_ = exports.download(self.store,created["export_id"],self.user)
        self.assertEqual(before,after)

    def test_unknown_consent_private_sheet_and_tampered_pin_cannot_export(self):
        changed = self.request.members[0].model_copy(update={"sheet":self.request.members[0].sheet.model_copy(update={"hash":"0"*64})})
        with self.assertRaises(HTTPException):self.create(self.request.model_copy(update={"members":[changed]}))
        with self.store.connect(write=True) as db:db.execute("UPDATE cases SET consent_confirmed=0 WHERE case_id=?",(self.case_id,))
        with self.assertRaises(HTTPException):self.create()
        with self.store.connect(write=True) as db:db.execute("UPDATE cases SET consent_confirmed=1 WHERE case_id=?",(self.case_id,))
        with self.assertRaises(HTTPException):exports.create(self.store,self.request,self.fixture.helper.people["admin"])

    def test_late_deletion_blocks_adoption(self):
        render = exports.render
        def deleted(snapshot):
            data = render(snapshot)
            with self.store.connect(write=True) as db:db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?",(self.case_id,))
            return data
        with patch.object(exports,"render",side_effect=deleted),self.assertRaises(HTTPException):self.create()
        with self.store.connect() as db:self.assertEqual(db.execute("SELECT COUNT(*) FROM changes WHERE action=?",(exports.ACTION,)).fetchone()[0],0)

    def synthetic_ai_pair(self):
        """Synthetic immutable AI row plus real grant/reveal; no provider invocation."""
        human = self.fixture.sheet
        data = sheets.view(self.store,human["sheet_id"],self.user)["document"]
        data.update(assigned_username=AI_ACCOUNT,origin="ai_service",ai_run_id="synthetic-ai-run",previous=[],initial_submission=None,revision=1)
        data["sheet"].update(sheet_id=uid(),rater_id="synthetic-ai",rater_kind="ai")
        doc = SheetDocumentV4.model_validate_json(encode(data))
        ref,digest = sheets.write_document(self.store,data)
        with self.store.connect(write=True) as db:
            db.execute("INSERT OR IGNORE INTO users(username,role,password_hash,active) VALUES(?,?,?,0)",(AI_ACCOUNT,"developer","service-only"))
            db.execute("INSERT INTO score_sheets VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",(doc.sheet.sheet_id,self.case_id,self.session_id,
                AI_ACCOUNT,"synthetic-ai",doc.rater_name,doc.source_hash,"independent","submitted",1,1,ref,digest))
        target = {"sheet_id":doc.sheet.sheet_id,"revision":1,"manifest_ref":ref,"manifest_hash":digest}
        self.fixture.helper.grant(human,target)
        exposed = self.fixture.helper.reveal(human,target,"operator")
        selection = model(exports.ExportSelectionV4,case_id=self.case_id,session_id=self.session_id,
            sheet={"sheet_id":doc.sheet.sheet_id,"revision":1,"ref":ref,"hash":digest},viewer_sheet_id=human["sheet_id"])
        return self.request.model_copy(update={"members":[self.request.members[0],selection]}),exposed

    def test_independent_historical_human_and_initial_ai_have_item_denominators(self):
        request,exposed = self.synthetic_ai_pair()
        created = self.create(request)
        snapshot = self.snapshot(created["export_id"])
        values = exports.tables(snapshot)
        valid = [row for row in values["independent_pairs"] if row["included"]]
        self.assertEqual([row["code"] for row in valid],["보23"])
        self.assertEqual(next(row for row in values["item_denominators"] if row["code"] == "보23")["valid_pair_n"],1)
        self.assertEqual(len(values["independent_pairs"]),83)
        self.assertNotIn("overall_accuracy",values)
        human = exposed["viewer"]
        selected = model(exports.ExportSelectionV4,case_id=self.case_id,session_id=self.session_id,
            sheet={"sheet_id":human["sheet_id"],"revision":human["revision"],"ref":human["manifest_ref"],"hash":human["manifest_hash"]})
        review = self.create(request.model_copy(update={"request_id":"exposed-review","members":[selected,request.members[1]]}))
        self.assertFalse(any(row["included"] for row in exports.tables(self.snapshot(review["export_id"]))["independent_pairs"]))
        with self.store.connect(write=True) as db:
            db.execute("DELETE FROM score_grants WHERE viewer_sheet_id=?",(human["sheet_id"],))
        with self.assertRaises(HTTPException):exports.download(self.store,created["export_id"],self.user)

    def test_download_detects_output_parent_and_late_receipt_revocation(self):
        created = self.create(); snapshot = self.snapshot(created["export_id"])
        for file in (snapshot.output,snapshot.parent_files[0]):
            path = self.store.path(file.ref); original = path.read_bytes(); path.write_bytes(original+b"changed")
            with self.assertRaises(HTTPException):exports.download(self.store,created["export_id"],self.user)
            path.write_bytes(original)
        read = exports.reports._read_file
        def revoke(store,pointer,stamps=None):
            data = read(store,pointer,stamps)
            if pointer.ref == snapshot.output.ref:
                with store.connect(write=True) as db:db.execute("UPDATE upload_receipts SET state='failed'")
            return data
        with patch.object(exports.reports,"_read_file",side_effect=revoke),self.assertRaises(HTTPException):
            exports.download(self.store,created["export_id"],self.user)

    def test_backup_restore_and_deleted_export_file_closure(self):
        created = self.create(); root = Path(self.fixture.helper.temp.name)
        expected = exports.download(self.store,created["export_id"],self.user)[0]
        maintenance.backup(self.store,root/"export-backup","operator")
        maintenance.restore(self.store,root/"export-backup",root/"export-restored")
        restored = Store(root/"export-restored")
        self.assertEqual(exports.download(restored,created["export_id"],self.user)[0],expected)
        with restored.connect(write=True) as db:
            db.execute("UPDATE cases SET deletion_requested=1 WHERE case_id=?",(self.case_id,))
        maintenance.clean(restored,purge_deleted=True)
        self.assertFalse(any((root/"export-restored"/"exports"/created["export_id"]).rglob("*.*")))

    def test_real_raw_zero_negative_and_f_g_memo_walk_tables_remain_separate(self):
        row = sheets.revise(self.store,self.fixture.sheet["sheet_id"],
            model(sheets.SheetReasonV4,expected_revision=self.fixture.sheet["revision"],reason="synthetic correction"),self.user,"reopen")
        with self.store.connect() as db:name = self.store.case(db,self.case_id)["dog_name"]
        for value in (-2,0):
            observation = self.fixture.helper.observe(value=value)
            observation["evidence"][0]["note"] = name+" F observed"
            observation["review_memo"] = "\t=G review "+name
            row = self.fixture.helper.save(row,[observation],"operator",
                linked_memos=[{"text":name+" linked memo","item_codes":["개5"],"evidence":observation["evidence"]}],
                walk_phases=[{"code":"개38","proximity_exception":"recheck","note":"synthetic walking recheck"}])
            selection = model(exports.ExportSelectionV4,case_id=self.case_id,session_id=self.session_id,
                sheet={"sheet_id":row["sheet_id"],"revision":row["revision"],"ref":row["manifest_ref"],"hash":row["manifest_hash"]})
            created = self.create(self.request.model_copy(update={"request_id":"raw-"+str(value),"members":[selection]}))
            tables = self.csv_tables(exports.download(self.store,created["export_id"],self.user)[0])
            raw = next(item for item in tables["raw_observations"] if item["code"] == "개5")
            self.assertEqual(raw["value"],str(value))
            self.assertIn("F observed",raw["evidence"])
            self.assertTrue(raw["review_memo"].startswith("'\t=G"))
            self.assertNotIn(name,raw["evidence"]+raw["review_memo"])
            self.assertEqual(tables["walk_exceptions"][0]["proximity_exception"],"recheck")
            self.assertEqual(tables["linked_memos"][0]["code"],"개59")
            self.assertTrue(tables["linked_memos"][0]["evidence"])
            self.assertFalse(tables["automatic_calculations"])

    def test_three_views_are_one_pair_and_different_batch_has_no_valid_denominator(self):
        request,_ = self.synthetic_ai_pair()
        snapshot = exports._snapshot(self.store,request,self.user)
        copied = []
        for member in snapshot.members:
            rows = []
            for row in member.sheet.sheet.observations:
                if row.code == "보23":
                    evidence = row.evidence[0]
                    row = row.model_copy(update={"evidence":tuple(evidence.model_copy(update={"camera_id":camera}) for camera in ("CAM1","CAM2","CAM3"))})
                rows.append(row)
            copied.append(member.model_copy(update={"sheet":member.sheet.model_copy(update={"sheet":member.sheet.sheet.model_copy(update={"observations":tuple(rows)})})}))
        snapshot = snapshot.model_copy(update={"members":tuple(copied)})
        values = exports.tables(snapshot)
        self.assertEqual(next(row for row in values["item_denominators"] if row["code"] == "보23")["valid_pair_n"],1)
        observed = [row for row in values["raw_observations"] if row["code"] == "보23"]
        self.assertEqual(len(observed),2)
        self.assertEqual(len(json.loads(observed[0]["evidence"])),3)
        changed = copied[1].model_copy(update={"sheet":copied[1].sheet.model_copy(update={"source_hash":"0"*64})})
        different = exports.tables(snapshot.model_copy(update={"members":(copied[0],changed)}))
        self.assertTrue(all(row["valid_pair_n"] == 0 for row in different["item_denominators"]))
        self.assertTrue(all("different_input_or_batch" in row["excluded_reasons"] for row in different["independent_pairs"]))

    def test_historical_identity_in_free_text_is_redacted_after_name_and_id_change(self):
        old = {"dog_name":"SyntheticOldDog","guardian_name":"SyntheticOldGuardian","participant_id":"OLD-ID-123"}
        new = {"dog_name":"SyntheticNewDog","guardian_name":"SyntheticNewGuardian","participant_id":"NEW-ID-456"}
        row = sheets.revise(self.store,self.fixture.sheet["sheet_id"],
            model(sheets.SheetReasonV4,expected_revision=self.fixture.sheet["revision"],reason="synthetic historical note"),self.user,"reopen")
        observation = self.fixture.helper.observe()
        observation["evidence"][0]["note"] = " / ".join(old.values())
        observation["review_memo"] = " / ".join(old.values())
        row = self.fixture.helper.save(row,[observation],"operator")
        with self.store.connect(write=True) as db:
            case = self.store.case(db,self.case_id); manifest = self.store.manifest(case)
            manifest.participant_id = new["participant_id"]
            self.store.save(db,case,manifest,self.user.username,"case.update")
            db.execute("UPDATE cases SET dog_name=?,guardian_name=?,participant_id=? WHERE case_id=?",(*new.values(),self.case_id))
            self.store.audit(db,self.user.username,self.case_id,"case.identity",{"before":old,"after":new})
        selection = model(exports.ExportSelectionV4,case_id=self.case_id,session_id=self.session_id,
            sheet={"sheet_id":row["sheet_id"],"revision":row["revision"],"ref":row["manifest_ref"],"hash":row["manifest_hash"]})
        created = self.create(self.request.model_copy(update={"members":[selection]}))
        values = self.csv_tables(exports.download(self.store,created["export_id"],self.user)[0])
        rendered = encode(values)
        self.assertTrue(all(name not in rendered for name in (*old.values(),*new.values())))
        raw = next(row for row in values["raw_observations"] if row["code"] == "개5")
        self.assertEqual(raw["review_memo"],"[비식별] / [비식별] / [비식별]")
