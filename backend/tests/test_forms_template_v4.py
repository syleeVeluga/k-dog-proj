"""U02 standard upload templates: headers, text IDs, validation lists and a filled file that previews unchanged."""

import base64
from io import BytesIO

from openpyxl import load_workbook

from app.api import create_app
from app.forms_v4 import TEMPLATE_PARTICIPANT, template, template_header
from app.domain.catalog import SURVEY_IDS
from tests.support import AppCase

ROSTER_HEADERS = ["참가자 ID", "반려견 이름", "보호자 이름", "예약 시각", "순번", "기본 동의", "분석·피드백 동의",
                  "낯선 사람 접촉 동의", "견종", "성별", "나이", "크기", "함께 산 기간", "입양 경로"]
SURVEY_HEADERS = [f"Q{index:02d}" for index in range(1, 29)]


class FormsTemplateV4Tests(AppCase):
    def setUp(self):
        super().setUp()
        self.app = create_app(self.root)
        self.client = self.client_for("operator")
        self.version = self.client.get("/api/catalog/survey").json()["version"]

    def test_headers_match_screen_item_names(self):
        self.assertEqual([label for _, label in TEMPLATE_PARTICIPANT], ROSTER_HEADERS)
        self.assertEqual([template_header(question) for question in SURVEY_IDS], SURVEY_HEADERS)

    def test_csv_has_bom_and_header_only(self):
        roster = template("roster", "csv")
        self.assertTrue(roster.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(roster.decode("utf-8-sig").strip().split(","), ROSTER_HEADERS)
        full = template("full", "csv").decode("utf-8-sig").strip().split(",")
        self.assertEqual(full, ROSTER_HEADERS + SURVEY_HEADERS)

    def test_xlsx_guides_first_keeps_ids_as_text_and_lists_choices(self):
        workbook = load_workbook(BytesIO(template("full", "xlsx")))
        self.assertEqual(workbook.sheetnames, ["안내", "입력"])
        self.assertEqual(workbook.active.title, "입력")
        self.assertIn("행사 ID는 파일이 아니라 업로드 화면에서", "\n".join(str(row[0].value) for row in workbook["안내"].iter_rows()))
        sheet = workbook["입력"]
        self.assertEqual([cell.value for cell in sheet[1]], ROSTER_HEADERS + SURVEY_HEADERS)
        self.assertEqual(sheet["A2"].number_format, "@")
        self.assertEqual(sheet["A1000"].number_format, "@")
        rules = {str(rule.sqref): rule for rule in sheet.data_validations.dataValidation}
        self.assertEqual(rules["J2:J1001"].formula1, '"암,수,중성화,미기재"')
        self.assertEqual(rules["K2:K1001"].type, "whole")
        # Q10~Q14 allow 0~4; the other questions 1~5.
        question = {name: rules[f"{sheet.cell(row=1, column=index).column_letter}2:{sheet.cell(row=1, column=index).column_letter}1001"]
                    for index, name in enumerate(ROSTER_HEADERS + SURVEY_HEADERS, 1) if name in ("Q01", "Q10")}
        self.assertEqual((question["Q01"].formula1, question["Q01"].formula2), ("1", "5"))
        self.assertEqual((question["Q10"].formula1, question["Q10"].formula2), ("0", "4"))
        self.assertEqual(len(load_workbook(BytesIO(template("roster", "xlsx")))["입력"][1]), len(ROSTER_HEADERS))

    def test_endpoint_roles_and_download_names(self):
        response = self.client.get("/api/forms/template?kind=roster&format=csv")
        self.assertEqual(response.status_code, 200)
        self.assertIn("kdog-roster.csv", response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        response = self.client.get("/api/forms/template?kind=full&format=xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("kdog-roster-survey.xlsx", response.headers["content-disposition"])
        for role, status in ((None, 401), ("reviewer", 403), ("developer", 403)):
            self.assertEqual(self.client_for(role).get("/api/forms/template?kind=roster").status_code, status)
        self.assertEqual(self.client.get("/api/forms/template?kind=other").status_code, 422)

    def test_filled_template_previews_without_changing_any_column(self):
        header = template("full", "csv").decode("utf-8-sig").strip().split(",")
        row = ["0001", "초코", "김보호", "", "1", "예", "확인", "미확인", "믹스", "수", "3", "소형", "2년", "입양"] + ["3"] * 28
        row[header.index("Q10")] = "0"
        source = ",".join(header) + "\n" + ",".join(row) + "\n"
        keys = [key for key, _ in TEMPLATE_PARTICIPANT] + list(SURVEY_IDS)
        columns = dict(zip(keys, header))
        config = {"request_id": "template-1", "event_id": "U02", "filename": "kdog-roster-survey.csv", "format": "csv",
                  "mapping": {"version": "forms-map-1", "source_profile": "google-forms", "survey_version": self.version,
                              "columns": columns, "answer_labels": {}}, "targets": []}
        response = self.client.post("/api/forms/preview", json={"config": config,
            "file_base64": base64.b64encode(source.encode()).decode()})
        self.assertEqual(response.status_code, 200, response.text)
        preview = response.json()
        self.assertEqual(preview["errors"], [])
        self.assertEqual(preview["rows"][0]["participant_id"], "0001")
        self.assertEqual(preview["rows"][0]["answers"]["s10"], 0)
        self.assertEqual(preview["rows"][0]["registration_status"], "all_answered")

    def test_filled_xlsx_opens_on_the_input_sheet_and_previews(self):
        workbook = load_workbook(BytesIO(template("roster", "xlsx")))
        sheet = workbook["입력"]
        sheet.append(["0007", "보리", "", "", 7, "", "", "", "", "", "", "", "", ""])
        output = BytesIO()
        workbook.save(output)
        found = self.client.post("/api/forms/columns?format=xlsx", content=output.getvalue())
        self.assertEqual(found.status_code, 200, found.text)
        self.assertEqual(found.json()["selected_sheet"], "입력")
        self.assertEqual([column["key"] for column in found.json()["columns"]], ROSTER_HEADERS)
        columns = {key: label for key, label in TEMPLATE_PARTICIPANT}
        config = {"request_id": "template-2", "event_id": "U02", "filename": "kdog-roster.xlsx", "format": "xlsx",
                  "mapping": {"version": "forms-map-1", "source_profile": "google-forms", "survey_version": self.version,
                              "columns": columns, "sheet": "입력", "answer_labels": {}}, "targets": []}
        response = self.client.post("/api/forms/preview", json={"config": config,
            "file_base64": base64.b64encode(output.getvalue()).decode()})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["errors"], [])
        self.assertEqual(response.json()["rows"][0]["participant_id"], "0007")
        self.assertEqual(response.json()["rows"][0]["registration_status"], "unregistered")
