"""S1 Forms file preview/commit with immutable provenance and atomic row application."""

import csv
from datetime import date, datetime
import hashlib
from io import BytesIO, StringIO
import json
import os
import sqlite3
from typing import Annotated, Literal, Self

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import Field, ValidationError, model_validator

from .domain.catalog import SURVEY_IDS
from .domain.catalog_v3 import SurveyCatalogV3
from .domain.catalog_v4 import RESOURCES, SURVEY_VERSION
from .input_models import Key, Model, Text
from .input_models_v3 import CaseCreateV3, SurveyEditV3
from .input_models_v4 import ManifestV4, new_session_v4
from .intake import participant_row, read_rows, selected_session
from .storage import encode, now, uid

Hash = Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]
Answer = Annotated[int, Field(ge=0, le=5)] | None
Cell = str | int | float | bool | None
PARTICIPANT_COLUMNS = {"event_id", "participant_id", "dog_name", "guardian_name", "reservation_at", "sequence_no",
                       "consent_confirmed", "consent_analysis_feedback", "consent_stranger_contact", "dog_breed",
                       "dog_sex", "dog_age_years", "dog_size", "years_together", "adoption_route"}
CANONICAL_COLUMNS = PARTICIPANT_COLUMNS | set(SURVEY_IDS) | {f"{q}_reason" for q in SURVEY_IDS}
MAX_BYTES = 32 * 1024 * 1024
# 표준 양식의 머리글은 화면 항목 이름과 같다. 행사 ID는 업로드 화면에서 받으므로 양식에 넣지 않는다.
TEMPLATE_PARTICIPANT = (
    ("participant_id", "참가자 ID"), ("dog_name", "반려견 이름"), ("guardian_name", "보호자 이름"),
    ("reservation_at", "예약 시각"), ("sequence_no", "순번"), ("consent_confirmed", "기본 동의"),
    ("consent_analysis_feedback", "분석·피드백 동의"), ("consent_stranger_contact", "낯선 사람 접촉 동의"),
    ("dog_breed", "견종"), ("dog_sex", "반려견 성별"), ("dog_age_years", "반려견 나이"), ("dog_size", "반려견 크기"),
    ("years_together", "함께 산 기간"), ("adoption_route", "입양 경로"),
)
TEMPLATE_TEXT_COLUMNS = {"participant_id", "dog_name", "guardian_name", "reservation_at", "dog_breed", "years_together"}
TEMPLATE_CHOICES = {
    "consent_confirmed": ("예",), "consent_analysis_feedback": ("확인", "미확인", "거절"),
    "consent_stranger_contact": ("확인", "미확인", "거절"), "dog_sex": ("암", "수", "중성화", "미기재"),
    "dog_size": ("소형", "중형", "대형", "미기재"), "dog_age_years": None, "adoption_route": ("분양", "입양", "기타", "미기재"),
}
TEMPLATE_ROWS = 1000


class FormsMappingV4(Model):
    version: Key
    source_profile: Text
    survey_version: Text
    columns: dict[str, Text]
    answer_labels: dict[str, dict[str, Annotated[int, Field(ge=0, le=5)]]] = Field(default_factory=dict)
    sheet: Text | None = None
    layout: Literal["rows", "transposed"] = "rows"
    historical_reference_only: bool = False

    @model_validator(mode="after")
    def explicit_mapping(self) -> Self:
        if not self.columns or set(self.columns) - CANONICAL_COLUMNS:
            raise ValueError("표준 참가자 열과 s01~s28만 연결하세요.")
        columns = list(self.columns.values())
        if len(set(columns)) != len(columns):
            raise ValueError("한 원본 열을 여러 표준 열에 중복 연결할 수 없습니다.")
        if set(self.answer_labels) - set(SURVEY_IDS):
            raise ValueError("응답 라벨 연결은 s01~s28 문항별로 지정하세요.")
        if (self.survey_version == SURVEY_VERSION) == self.historical_reference_only:
            raise ValueError("현재 설문 판본만 S1에 반영하며 다른 판본은 원자료 참고로만 등록합니다.")
        for labels in self.answer_labels.values():
            if any(not label.strip() for label in labels) or len({label.strip() for label in labels}) != len(labels):
                raise ValueError("빈 라벨이나 공백만 다른 모호한 라벨을 연결할 수 없습니다.")
        return self


class FormsTargetV4(Model):
    row_number: Annotated[int, Field(ge=2, le=10001)]
    case_id: Key
    session_id: Key
    expected_revision: Annotated[int, Field(ge=1)]


class FormsPreviewRequestV4(Model):
    request_id: Key
    event_id: Key
    filename: Text
    format: Literal["csv", "xlsx"]
    mapping: FormsMappingV4
    targets: list[FormsTargetV4] = Field(default_factory=list)

    @model_validator(mode="after")
    def unique_targets(self) -> Self:
        if len({target.row_number for target in self.targets}) != len(self.targets):
            raise ValueError("각 원본 행은 하나의 대상에만 연결하세요.")
        if len({target.case_id for target in self.targets}) != len(self.targets):
            raise ValueError("한 파일의 여러 행을 같은 참가자에 덮어쓸 수 없습니다.")
        return self


class FormsIssueV4(Model):
    row_number: int | None = None
    message: Text


class FormsPreviewRowV4(Model):
    row_number: int
    source_location: Text
    action: Literal["create", "update"]
    case_id: Key
    session_id: Key
    event_id: Key
    participant_id: Key
    dog_name: Text
    expected_revision: int | None
    participant: CaseCreateV3 | None = None
    answers: dict[str, Answer]
    previous_answers: dict[str, Answer]
    previous_blank_reasons: dict[str, str] = Field(default_factory=dict)
    blank_reasons: dict[str, str]
    changed_questions: list[str]
    raw_values: dict[str, Cell]
    registration_status: Literal["unregistered", "partial", "all_answered", "reference_only"]
    calculation_status: Literal["current_edition", "reference_only"]


class FormsPreviewV4(Model):
    schema_version: Literal["forms-preview-4.0"] = "forms-preview-4.0"
    preview_id: Key
    preview_hash: Hash
    source_hash: Hash
    filename: Text
    source_profile: Text
    mapping_version: Key
    survey_version: Text
    source_sheet: str | None
    original_headers: list[str]
    rows: list[FormsPreviewRowV4]
    errors: list[FormsIssueV4]
    already_committed: bool = False


class FormsCommitRequestV4(Model):
    request_id: Key
    preview_id: Key
    preview_hash: Hash


class FormsCommitRowV4(Model):
    row_number: int
    case_id: Key
    session_id: Key
    event_id: Key
    participant_id: Key
    input_revision: int
    registration_status: Literal["registered", "reference_only"]


class FormsCommitResultV4(Model):
    schema_version: Literal["forms-commit-4.0"] = "forms-commit-4.0"
    preview_id: Key
    rows: list[FormsCommitRowV4]


def _digest(value) -> str:
    return hashlib.sha256(value if isinstance(value, bytes) else encode(value).encode("utf-8")).hexdigest()


def _writer(db, actor: str) -> None:
    user = db.execute("SELECT role,active FROM users WHERE username=?", (actor,)).fetchone()
    if not user or not user["active"] or user["role"] not in ("operator", "admin"):
        raise HTTPException(403, "현재 활성 운영자 또는 관리자 권한이 필요합니다.")


def _find(db, action, target):
    row = db.execute("SELECT detail_json FROM changes WHERE action=? AND target=? LIMIT 1",
                     (action, target)).fetchone()
    return json.loads(row[0]) if row else None


def _request_target(actor, request_id):
    return _digest([actor, request_id])


def _check_request(db, action, actor, request_id, request_hash):
    receipt = _find(db, action, _request_target(actor, request_id))
    if receipt and receipt["request_hash"] != request_hash:
        raise HTTPException(409, "같은 요청 ID에 다른 파일 또는 설정을 사용할 수 없습니다.")
    return receipt


def _write(store, kind, data: bytes, suffix="json"):
    ref = f"inputs/forms-{kind}-{uid()}.{suffix}"
    with store.path(ref).open("xb") as handle:
        handle.write(data)
        handle.flush()
        os.fsync(handle.fileno())
    return {"ref": ref, "hash": _digest(data)}


def _read(store, link):
    try:
        data = store.path(link["ref"]).read_bytes()
    except (KeyError, OSError) as exc:
        raise HTTPException(409, "등록 원본 파일을 읽을 수 없습니다.") from exc
    if _digest(data) != link["hash"]:
        raise HTTPException(409, "등록 원본 파일 해시가 일치하지 않습니다.")
    return data


def _load_preview(store, db, preview_id, expected_hash=None):
    link = _find(db, "forms.preview", preview_id)
    if link is None:
        raise HTTPException(409, "미리보기 기록이 없거나 삭제 처리되었습니다. 새로 확인하세요.")
    if expected_hash is not None and link["hash"] != expected_hash:
        raise HTTPException(409, "미리보기 고정 해시가 다릅니다.")
    record = json.loads(_read(store, link))
    _read(store, record["source"])
    request = FormsPreviewRequestV4.model_validate(record["request"])
    result = FormsPreviewV4.model_validate({**record["view"], "preview_hash": link["hash"]})
    if result.preview_id != preview_id or result.source_hash != record["source"]["hash"]:
        raise HTTPException(409, "미리보기와 원파일 참조가 일치하지 않습니다.")
    return record, request, result


def _assert_live(store, db, preview, *, committed=False, associated_case_ids=(), associated_participants=()):
    case_ids = set(associated_case_ids)
    for identity in associated_participants:
        pair = (identity["event_id"], identity["participant_id"])
        current = db.execute("SELECT case_id FROM cases WHERE event_id=? AND participant_id=?", pair).fetchone()
        if current:
            case_ids.add(current["case_id"])
        if db.execute("SELECT 1 FROM changes WHERE action='deletion.record' AND json_extract(detail_json,'$.event_id')=? "
                      "AND json_extract(detail_json,'$.participant_id')=?", pair).fetchone():
            raise HTTPException(409, "삭제된 참가자의 미리보기는 재사용할 수 없습니다.")
    for case_id in case_ids:
        store.case(db, case_id)
    for item in preview.rows:
        if committed or item.action == "update":
            row = store.case(db, item.case_id, expected=None if committed else item.expected_revision)
            manifest = store.manifest(row)
            if not isinstance(manifest, ManifestV4):
                raise HTTPException(409, "S1 입력 전환을 완료한 참가자만 연결할 수 있습니다.")
            if not committed:
                selected_session(manifest, item.session_id)
            if (row["event_id"], row["participant_id"]) != (item.event_id, item.participant_id):
                raise HTTPException(409, "참가자 식별 정보가 변경되었습니다. 새 미리보기가 필요합니다.")


def _labels(mapping: FormsMappingV4, catalog: SurveyCatalogV3):
    output = {}
    for question in catalog.items:
        values = {str(value): value for value in question.allowed_values}
        values.update({label.text.strip(): label.value for label in question.labels})
        for label, value in mapping.answer_labels.get(question.item_id, {}).items():
            label = label.strip()
            if value not in question.allowed_values or label in values and values[label] != value:
                raise HTTPException(422, f"{question.item_id}: 척도와 충돌하거나 모호한 응답 라벨 연결입니다.")
            values[label] = value
        output[question.item_id] = values
    return output


def _scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or type(value) in (str, int, float, bool):
        return value
    raise ValueError("지원하지 않는 원본 셀 형식입니다.")


def _table(data, format, sheet=None, layout="rows"):
    if format not in ("csv", "xlsx") or layout not in ("rows", "transposed"):
        raise HTTPException(422, "지원하지 않는 파일 형식 또는 배치입니다.")
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(422, "32 MiB 이하의 비어 있지 않은 원파일이 필요합니다.")
    metadata = {}
    rows = read_rows(data, format, sheet, metadata)
    if not rows or len(rows) > 10001 or max((len(row) for row in rows), default=0) > 100:
        raise HTTPException(422, "헤더와 최대 10,000개 입력 행, 100개 열이 필요합니다.")
    if any(isinstance(cell, str) and cell.lstrip().startswith("=") for row in rows for cell in row):
        raise HTTPException(422, "수식 셀은 가져올 수 없습니다. 원응답으로 내보내세요.")
    original_rows = [[_scalar(value) for value in row] for row in rows]
    if layout == "transposed":
        if len({len(row) for row in rows}) != 1:
            raise HTTPException(422, "전치 프로필은 행과 열이 맞는 표가 필요합니다.")
        rows = [list(row) for row in zip(*rows)]
    if not rows or not all(isinstance(header, str) and header.strip() for header in rows[0]):
        raise HTTPException(422, "원본 헤더는 비어 있지 않은 문자열이어야 합니다.")
    headers = rows[0]
    if len({header.strip() for header in headers}) != len(headers):
        raise HTTPException(422, "중복되거나 공백만 다른 원본 헤더를 수정하세요.")
    return rows, metadata, original_rows


def template_header(question_id: str) -> str:
    return f"Q{question_id[1:]}"


def template(kind: Literal["roster", "full"], format: Literal["csv", "xlsx"]) -> bytes:
    """명단만 또는 명단+설문 28문항 표준 양식. 머리글은 화면 항목 이름이고 ID 열은 텍스트로 지정한다."""
    columns = [(key, label) for key, label in TEMPLATE_PARTICIPANT]
    ranges = {}
    if kind == "full":
        catalog = SurveyCatalogV3.model_validate_json((RESOURCES / "catalogs/survey-v3.json").read_bytes())
        for item in catalog.items:
            columns.append((item.item_id, template_header(item.item_id)))
            ranges[item.item_id] = (min(item.allowed_values), max(item.allowed_values))
    headers = [label for _, label in columns]
    if format == "csv":
        output = StringIO(newline="")
        csv.writer(output).writerow(headers)
        return output.getvalue().encode("utf-8-sig")
    workbook = Workbook()
    guide = workbook.active
    guide.title = "안내"
    lines = ["K-DOG 일괄 등록 양식 — " + ("명단 + 설문 28문항" if kind == "full" else "명단만"),
             "「입력」 시트에 참가자 한 명을 한 줄씩 적습니다. 1행(머리글)은 바꾸지 마세요.",
             "반려견 이름은 꼭 필요합니다. 참가자 ID는 비우면 새 ID가 발급되어 고친 파일을 다시 올릴 때 중복 등록되므로 꼭 적으세요. 나머지는 모르면 비워 둡니다.",
             "참가자 ID는 텍스트로 입력됩니다. 0001처럼 앞자리 0까지 그대로 적으세요.",
             "행사 ID는 파일이 아니라 업로드 화면에서 입력합니다.",
             "기본 동의는 동의서를 확인했을 때만 「예」로 적고, 아니면 비웁니다. 나머지 동의는 확인·미확인·거절 중 하나입니다.",
             "연락처는 받지 않습니다."]
    if kind == "full":
        lines.append("설문 Q01~Q28은 응답 숫자를 적습니다. Q10~Q14는 0~4, 나머지는 1~5입니다. 응답하지 않은 문항은 비웁니다. "
                     "빈칸 사유를 남기려면 입력 시트 오른쪽에 「Q01 빈칸 사유」처럼 열을 추가하세요.")
    lines.append("Excel에서 CSV로 저장하려면 「CSV UTF-8(쉼표로 분리)」를 고르세요.")
    for line in lines:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    sheet = workbook.create_sheet("입력")
    sheet.append(headers)
    last = TEMPLATE_ROWS + 1
    for index, (key, _) in enumerate(columns, 1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = max(12, len(headers[index - 1]) * 2 + 2)
        if key in TEMPLATE_TEXT_COLUMNS:
            for row in range(2, last + 1):
                sheet.cell(row=row, column=index).number_format = "@"
        if key in TEMPLATE_CHOICES and TEMPLATE_CHOICES[key]:
            validation = DataValidation(type="list", formula1='"' + ",".join(TEMPLATE_CHOICES[key]) + '"', allow_blank=True)
            validation.error = "목록에 있는 값만 쓸 수 있습니다: " + ", ".join(TEMPLATE_CHOICES[key])
        elif key == "dog_age_years":
            validation = DataValidation(type="whole", operator="between", formula1="0", formula2="30", allow_blank=True)
            validation.error = "0~30 사이의 정수(세)만 쓸 수 있습니다. 모르면 비우세요."
        elif key == "sequence_no":
            validation = DataValidation(type="whole", operator="between", formula1="1", formula2="9999", allow_blank=True)
            validation.error = "1~9999 사이의 정수만 쓸 수 있습니다."
        elif key in ranges:
            low, high = ranges[key]
            validation = DataValidation(type="whole", operator="between", formula1=str(low), formula2=str(high), allow_blank=True)
            validation.error = f"{low}~{high} 사이의 정수만 쓸 수 있습니다. 응답하지 않았으면 비우세요."
        else:
            continue
        validation.errorTitle = "입력할 수 없는 값"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last}")
    sheet.freeze_panes = "A2"
    # 업로드 화면은 활성 시트를 기본으로 읽으므로 안내 시트가 아니라 입력 시트를 활성으로 둔다.
    guide.sheet_view.tabSelected = False
    sheet.sheet_view.tabSelected = True
    workbook.active = 1
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def inspect_columns(data: bytes, format: str, sheet: str | None = None, layout: str = "rows") -> dict:
    rows, metadata, _ = _table(data, format, sheet, layout)
    return {"columns": [{"key": header, "label": header} for header in rows[0]],
            "sheets": metadata.get("sheets", []), "selected_sheet": metadata.get("selected_sheet")}


def _case_associations(store, db, rows, request):
    # Error rows still hold original personal data. Associate by explicit IDs, never names.
    case_ids = {target.case_id for target in request.targets}
    participants = set()
    headers = rows[0]
    for cells in rows[1:]:
        raw = dict(zip(headers, cells))
        participant_id = raw.get(request.mapping.columns.get("participant_id"))
        event_id = raw.get(request.mapping.columns.get("event_id")) or request.event_id
        if not isinstance(participant_id, str) or not isinstance(event_id, str):
            continue
        participants.add((event_id, participant_id))
        current = db.execute("SELECT case_id FROM cases WHERE event_id=? AND participant_id=?", (event_id, participant_id)).fetchone()
        if current:
            case_ids.add(current["case_id"])
        if db.execute("SELECT 1 FROM changes WHERE action='deletion.record' AND json_extract(detail_json,'$.event_id')=? "
                      "AND json_extract(detail_json,'$.participant_id')=?", (event_id, participant_id)).fetchone():
            raise HTTPException(409, "삭제 이력이 있는 원행은 가져올 수 없습니다.")
    for case_id in case_ids:
        # No new preview/source blob may retain a participant whose deletion is already requested.
        store.case(db, case_id)
    return sorted(case_ids), [{"event_id": event, "participant_id": participant} for event, participant in sorted(participants)]


def _build_rows(store, db, rows, request, catalog, sheet_name):
    mapping = request.mapping
    labels = _labels(mapping, catalog) if not mapping.historical_reference_only else {}
    targets = {target.row_number: target for target in request.targets}
    planned, errors, identities, sequences, encountered = [], [], set(), set(), set()
    headers = rows[0]
    for number, cells in enumerate(rows[1:], 2):
        if not any(value not in (None, "") for value in cells):
            continue
        encountered.add(number)
        try:
            if len(cells) != len(headers):
                raise ValueError("열 수가 헤더와 다릅니다.")
            original = {header: _scalar(value) for header, value in zip(headers, cells)}
            mapped = {key: original[header] for key, header in mapping.columns.items()}
            if mapped.get("event_id") not in (None, "", request.event_id):
                raise ValueError("원행의 행사 ID가 선택한 행사와 다릅니다.")
            target = targets.get(number)
            before, previous_reasons = dict.fromkeys(SURVEY_IDS), {}
            if target:
                current = store.case(db, target.case_id, expected=target.expected_revision)
                manifest = store.manifest(current)
                if not isinstance(manifest, ManifestV4) or current["event_id"] != request.event_id:
                    raise ValueError("선택한 행사와 S1 참가자에 연결하세요.")
                session = selected_session(manifest, target.session_id)
                if mapped.get("participant_id") not in (None, "", current["participant_id"]):
                    raise ValueError("원행의 참가자 ID가 명시적으로 연결한 대상과 다릅니다.")
                if not mapping.historical_reference_only and session.survey_version != mapping.survey_version:
                    raise ValueError("대상 회차의 원설문 판본과 일치하지 않습니다.")
                participant = None
                case_id, session_id, participant_id = current["case_id"], session.session_id, current["participant_id"]
                dog_name, revision, action = current["dog_name"], current["input_revision"], "update"
                before = session.survey.copy()
                previous_reasons = session.survey_blank_reasons.copy()
            else:
                participant_id = mapped.get("participant_id") or "F_" + uid()
                if not isinstance(participant_id, str):
                    raise ValueError("참가자 ID는 텍스트여야 합니다. 숫자 ID를 자동 변환하지 않습니다.")
                participant = participant_row({**mapped, "event_id": request.event_id, "participant_id": participant_id,
                                               "dog_name": mapped.get("dog_name")})
                if db.execute("SELECT 1 FROM cases WHERE event_id=? AND participant_id=?", (request.event_id, participant_id)).fetchone():
                    raise ValueError("이미 등록된 ID입니다. 기존 참가자와 회차를 명시적으로 연결하세요.")
                deleted = db.execute("SELECT 1 FROM changes WHERE action='deletion.record' AND json_extract(detail_json,'$.event_id')=? "
                                     "AND json_extract(detail_json,'$.participant_id')=?", (request.event_id, participant_id)).fetchone()
                if deleted:
                    raise ValueError("삭제 이력이 있는 참가자 ID는 가져올 수 없습니다.")
                case_id, session_id, dog_name, revision, action = uid(), uid(), participant.dog_name, None, "create"
                if participant.sequence_no is not None:
                    key = (request.event_id, participant.sequence_no)
                    if key in sequences or db.execute("SELECT 1 FROM cases WHERE event_id=? AND sequence_no=?", key).fetchone():
                        raise ValueError("행사 내 참가 순번이 중복됩니다.")
                    sequences.add(key)
            pair = (request.event_id, participant_id)
            if pair in identities:
                raise ValueError("파일 안의 참가자 ID가 중복됩니다.")
            identities.add(pair)
            answers, reasons = dict.fromkeys(SURVEY_IDS), {}
            if not mapping.historical_reference_only:
                for question in SURVEY_IDS:
                    raw = mapped.get(question)
                    if raw is None or raw == "":
                        value = None
                    elif type(raw) is int:
                        value = labels[question].get(str(raw))
                        if value is None:
                            raise ValueError(f"{question}: 허용된 척도 범위를 확인하세요.")
                    elif isinstance(raw, str) and raw.strip() in labels[question]:
                        value = labels[question][raw.strip()]
                    else:
                        raise ValueError(f"{question}: 허용 숫자 또는 명시한 응답 라벨이 아닙니다.")
                    answers[question] = value
                    reason = mapped.get(f"{question}_reason")
                    if reason not in (None, ""):
                        if not isinstance(reason, str) or not reason.strip() or value is not None:
                            raise ValueError(f"{question}: 빈칸 사유는 실제 결측에만 기록하세요.")
                        reasons[question] = reason.strip()
                SurveyEditV3(expected_revision=revision or 1, session_id=session_id,
                             survey_version=SURVEY_VERSION, answers=answers, blank_reasons=reasons)
            count = sum(value is not None for value in answers.values())
            planned.append(FormsPreviewRowV4(row_number=number,
                source_location=f"{sheet_name or 'CSV'} {'열' if mapping.layout == 'transposed' else '행'} {number}",
                action=action, case_id=case_id, session_id=session_id, event_id=request.event_id,
                participant_id=participant_id, dog_name=dog_name, expected_revision=revision, participant=participant,
                answers=answers, previous_answers=before, previous_blank_reasons=previous_reasons, blank_reasons=reasons,
                changed_questions=[] if mapping.historical_reference_only else [q for q in SURVEY_IDS
                    if before[q] != answers[q] or previous_reasons.get(q) != reasons.get(q)],
                raw_values=original, registration_status="reference_only" if mapping.historical_reference_only else
                    "all_answered" if count == 28 else "partial" if count or reasons else "unregistered",
                calculation_status="reference_only" if mapping.historical_reference_only else "current_edition"))
        except (ValueError, HTTPException) as exc:
            message = ("; ".join(".".join(map(str, error["loc"])) + ": " + error["msg"] for error in exc.errors())
                       if isinstance(exc, ValidationError) else str(exc.detail) if isinstance(exc, HTTPException) else str(exc))
            errors.append(FormsIssueV4(row_number=number, message=message[:200]))
    for missing in set(targets) - encountered:
        errors.append(FormsIssueV4(row_number=missing, message="연결 대상에 지정한 원본 행이 없거나 비어 있습니다."))
    if not planned and not errors:
        errors.append(FormsIssueV4(message="등록할 응답 행이 없습니다."))
    return planned, errors


def preview(store, data: bytes, request: FormsPreviewRequestV4, actor: str) -> FormsPreviewV4:
    request = FormsPreviewRequestV4.model_validate_json(request.model_dump_json())
    source_hash = _digest(data)
    request_hash = _digest({"request": request.model_dump(mode="json"), "source_hash": source_hash})
    fingerprint = _digest({"event_id": request.event_id, "mapping": request.mapping.model_dump(mode="json"),
                           "targets": [t.model_dump(mode="json") for t in request.targets], "source_hash": source_hash})
    with store.connect(write=True) as db:
        _writer(db, actor)
        receipt = _check_request(db, "forms.preview_request", actor, request.request_id, request_hash)
        if db.execute("SELECT 1 FROM changes WHERE action='forms.purged' AND json_extract(detail_json,'$.source_hash')=?",
                      (source_hash,)).fetchone():
            raise HTTPException(409, "삭제 대상이 포함된 가져오기 파일은 다시 등록할 수 없습니다.")
        existing = _find(db, "forms.fingerprint", fingerprint)
        if receipt or existing:
            preview_id = (receipt or existing)["preview_id"]
            record, _, result = _load_preview(store, db, preview_id)
            committed = _find(db, "forms.commit", preview_id) is not None
            _assert_live(store, db, result, committed=committed, associated_case_ids=record.get("associated_case_ids", []),
                         associated_participants=record.get("associated_participants", []))
            if not receipt:
                store.audit(db, actor, _request_target(actor, request.request_id), "forms.preview_request",
                            {"request_hash": request_hash, "preview_id": preview_id})
            result.already_committed = committed
            return result
        rows, metadata, original_rows = _table(data, request.format, request.mapping.sheet, request.mapping.layout)
        if not set(request.mapping.columns.values()) <= set(rows[0]):
            raise HTTPException(422, "연결한 원본 헤더가 선택한 시트에 없습니다.")
        sheet_name = metadata.get("selected_sheet")
        associated_case_ids, associated_participants = _case_associations(store, db, rows, request)
        catalog = SurveyCatalogV3.model_validate_json((RESOURCES / "catalogs/survey-v3.json").read_bytes())
        plans, errors = _build_rows(store, db, rows, request, catalog, sheet_name)
        source = _write(store, "source", data, request.format)
        preview_id = uid()
        view = dict(preview_id=preview_id, source_hash=source_hash, filename=request.filename,
                    source_profile=request.mapping.source_profile, mapping_version=request.mapping.version,
                    survey_version=request.mapping.survey_version, source_sheet=sheet_name, original_headers=rows[0],
                    rows=[row.model_dump(mode="json") for row in plans], errors=[issue.model_dump() for issue in errors])
        record = {"schema_version": "forms-record-4.0", "request": request.model_dump(mode="json"),
                  "fingerprint": fingerprint, "source": source, "original_rows": original_rows,
                  "associated_case_ids": associated_case_ids, "associated_participants": associated_participants, "view": view}
        link = _write(store, "preview", encode(record).encode("utf-8"))
        store.audit(db, actor, preview_id, "forms.preview", {**link, "fingerprint": fingerprint, "source_hash": source_hash})
        store.audit(db, actor, fingerprint, "forms.fingerprint", {"preview_id": preview_id})
        store.audit(db, actor, _request_target(actor, request.request_id), "forms.preview_request",
                    {"request_hash": request_hash, "preview_id": preview_id})
        return FormsPreviewV4.model_validate({**view, "preview_hash": link["hash"]})


def _create_case(store, db, row, source_request, actor):
    participant = row.participant
    if participant is None:
        raise HTTPException(409, "신규 참가자의 고정 정보가 없습니다.")
    session = new_session_v4(row.session_id)
    if not source_request.mapping.historical_reference_only:
        session.survey, session.survey_blank_reasons = row.answers, row.blank_reasons
    manifest = ManifestV4(case_id=row.case_id, event_id=row.event_id, participant_id=row.participant_id,
                          input_revision=1, selected_session_id=row.session_id, sessions=[session], consents=participant.consents)
    ref, digest = store.write_manifest(manifest)
    db.execute("INSERT INTO cases(case_id,event_id,participant_id,dog_name,reservation_at,sequence_no,consent_confirmed,"
               "guardian_name,dog_profile_json,input_revision,selected_session_id,manifest_ref,manifest_hash,created_at,"
               "updated_at,manifest_schema_version,consents_v3_json) VALUES (?,?,?,?,?,?,?,?,?,1,?,?,?,?,?,?,?)",
               (row.case_id, row.event_id, row.participant_id, participant.dog_name, participant.reservation_at,
                participant.sequence_no, int(participant.consent_confirmed), participant.guardian_name,
                participant.dog.model_dump_json(), row.session_id, ref, digest, now(), now(), manifest.schema_version,
                manifest.consents.model_dump_json()))
    store.audit(db, actor, row.case_id, "case.create", {"revision": 1, "source": "forms"})
    return 1


def commit(store, request: FormsCommitRequestV4, actor: str) -> FormsCommitResultV4:
    request = FormsCommitRequestV4.model_validate_json(request.model_dump_json())
    request_hash = _digest(request.model_dump(mode="json"))
    with store.connect(write=True) as db:
        _writer(db, actor)
        _check_request(db, "forms.commit_request", actor, request.request_id, request_hash)
        record, source_request, frozen = _load_preview(store, db, request.preview_id, request.preview_hash)
        completed = _find(db, "forms.commit", request.preview_id)
        _assert_live(store, db, frozen, committed=completed is not None,
                     associated_case_ids=record.get("associated_case_ids", []),
                     associated_participants=record.get("associated_participants", []))
        if completed:
            result = FormsCommitResultV4.model_validate_json(_read(store, completed))
            if not _find(db, "forms.commit_request", _request_target(actor, request.request_id)):
                store.audit(db, actor, _request_target(actor, request.request_id), "forms.commit_request",
                            {"request_hash": request_hash, "preview_id": request.preview_id})
            return result
        if frozen.errors or not frozen.rows:
            raise HTTPException(422, "미리보기 오류를 모두 수정해야 전체 파일을 확정할 수 있습니다.")
        # Validate every row and existing state before the first participant write.
        for row in frozen.rows:
            if row.action == "create":
                participant = row.participant
                if (db.execute("SELECT 1 FROM cases WHERE case_id=? OR (event_id=? AND participant_id=?)",
                               (row.case_id, row.event_id, row.participant_id)).fetchone()
                        or participant.sequence_no is not None and db.execute(
                            "SELECT 1 FROM cases WHERE event_id=? AND sequence_no=?", (row.event_id, participant.sequence_no)).fetchone()):
                    raise HTTPException(409, "미리보기 이후 참가자 ID 또는 순번이 등록되었습니다. 다시 확인하세요.")
                if db.execute("SELECT 1 FROM changes WHERE action='deletion.record' AND json_extract(detail_json,'$.event_id')=? "
                              "AND json_extract(detail_json,'$.participant_id')=?", (row.event_id, row.participant_id)).fetchone():
                    raise HTTPException(409, "미리보기 이후 삭제 이력이 기록되었습니다.")
        results = []
        try:
            for row in frozen.rows:
                if row.action == "create":
                    revision = _create_case(store, db, row, source_request, actor)
                else:
                    current = store.case(db, row.case_id, expected=row.expected_revision)
                    manifest = store.manifest(current)
                    session = selected_session(manifest, row.session_id)
                    if not source_request.mapping.historical_reference_only:
                        if session.survey_version != source_request.mapping.survey_version:
                            raise HTTPException(409, "대상 회차의 설문 판본이 변경되었습니다.")
                        session.survey = row.answers
                        session.survey_not_applicable = []
                        session.survey_blank_reasons = row.blank_reasons
                    store.save(db, current, manifest, actor, "forms.survey.update")
                    revision = manifest.input_revision
                capsule = {"schema_version": "forms-source-row-4.0", "preview_id": request.preview_id,
                           "case_id": row.case_id, "session_id": row.session_id, "source_hash": frozen.source_hash,
                           "source_filename": frozen.filename, "source_location": row.source_location,
                           "mapping": source_request.mapping.model_dump(mode="json"), "raw_values": row.raw_values,
                           "registration_status": row.registration_status, "answers": row.answers,
                           "previous_answers": row.previous_answers, "previous_blank_reasons": row.previous_blank_reasons,
                           "blank_reasons": row.blank_reasons}
                link = _write(store, "row", encode(capsule).encode("utf-8"))
                store.audit(db, actor, row.case_id, "forms.source_row", {**link, "preview_id": request.preview_id,
                            "session_id": row.session_id, "input_revision": revision, "source_hash": frozen.source_hash})
                results.append(FormsCommitRowV4(row_number=row.row_number, case_id=row.case_id, session_id=row.session_id,
                    event_id=row.event_id, participant_id=row.participant_id, input_revision=revision,
                    registration_status="reference_only" if source_request.mapping.historical_reference_only else "registered"))
        except sqlite3.IntegrityError as exc:
            raise HTTPException(409, "등록 중 참가자 ID·순번 충돌이 발생했습니다. 전체 작업을 취소했습니다.") from exc
        result = FormsCommitResultV4(preview_id=request.preview_id, rows=results)
        link = _write(store, "receipt", result.model_dump_json().encode("utf-8"))
        store.audit(db, actor, request.preview_id, "forms.commit", link)
        store.audit(db, actor, _request_target(actor, request.request_id), "forms.commit_request",
                    {"request_hash": request_hash, "preview_id": request.preview_id})
        return result


def purge_case_imports(store, db, case_id: str, *, identity: dict | None = None) -> None:
    """Called by maintenance before purging a case; retain other participants' row capsules."""
    imports, affected_sources = [], set()
    case = db.execute("SELECT event_id,participant_id FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if identity is None and case:
        identity = {"event_id": case["event_id"], "participant_id": case["participant_id"]}
    for audit in db.execute("SELECT target,detail_json FROM changes WHERE action='forms.preview'").fetchall():
        link = json.loads(audit["detail_json"])
        record = json.loads(_read(store, link))
        source_hash = record["view"]["source_hash"]
        imports.append((audit["target"], record["fingerprint"], source_hash))
        if (case_id in record.get("associated_case_ids", [])
                or identity is not None and identity in record.get("associated_participants", [])
                or identity is not None and any((row["event_id"], row["participant_id"]) ==
                    (identity["event_id"], identity["participant_id"]) for row in record["view"]["rows"])
                or any(row["case_id"] == case_id for row in record["view"]["rows"])):
            affected_sources.add(source_hash)
    actor = db.execute("SELECT username FROM users ORDER BY username LIMIT 1").fetchone()
    for preview_id, fingerprint, source_hash in imports:
        if source_hash not in affected_sources:
            continue
        # Tombstones contain no rows, participant names, answers or live blob references.
        store.audit(db, actor[0], fingerprint, "forms.purged", {"source_hash": source_hash})
        db.execute("DELETE FROM changes WHERE target=? AND action IN ('forms.preview','forms.commit')", (preview_id,))
        db.execute("DELETE FROM changes WHERE target=? AND action='forms.fingerprint'", (fingerprint,))
