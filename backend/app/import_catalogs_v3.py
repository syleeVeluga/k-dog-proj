"""Read-only v3 extraction; formulas are preserved as source, never executed."""

import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .domain.catalog_v3 import (
    AUTO_CODES, UNUSED_CODES, SHEET_ROWS, SEGMENTS, BehaviorCatalogV3, MappingV3,
    ProtocolV3, ScoringRulesV3, SurveyCatalogV3,
)
from .domain.validation_v3 import validate_references_v3

ROOT = Path(__file__).resolve().parents[2]
CUSTOMER_DIR_V3 = ROOT / "docs/큐브전달_20260929"
WORKBOOK = "기준자료/02_개_행동채점표_최종정리_20260929.xlsx"
SURVEY = "기준자료/01_보호자_설문지_20260919.docx"
PROCEDURE = "기준자료/04_행동실험_진행절차_20260929.docx"
DECISIONS = "기준자료/07_결과판정_애착유형_사회성_행동동조_20260929.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source(root: Path, filename: str, location: str) -> dict:
    return {"filename": filename, "sha256": sha256(root / filename), "location": location}


def verify_attachments(root: Path) -> dict:
    manifest = json.loads((root / "첨부목록.json").read_text(encoding="utf-8"))
    for entry in manifest["files"]:
        path = (root / entry["file"]).resolve()
        if not path.is_relative_to(root.resolve()):
            raise ValueError("attachment path outside source folder")
        if path.stat().st_size != entry["bytes"] or sha256(path) != entry["sha256"]:
            raise ValueError(f"attachment size/hash differs: {entry['file']}")
    return manifest


def docx_paragraphs(root: Path, filename: str) -> list[dict]:
    with ZipFile(root / filename) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
    return [{"text": "".join(node.text or "" for node in paragraph.findall(".//w:t", NS)),
             "source": source(root, filename, f"word/document.xml paragraph {index}")}
            for index, paragraph in enumerate(tree.findall(".//w:p", NS), 1)]


def protocol_v3(root: Path) -> ProtocolV3:
    src = source(root, PROCEDURE, "§1·4 구간별 진행")
    windows = [{"window_id": segment, "segments": [segment], "anchor": "actual_segment",
                "actual_event_required": True, "source": src} for segment, _, _ in SEGMENTS]
    # Fixed offsets are relative to recorded segment/event anchors, never absolute video timestamps.
    timed = (
        ("alone_initial", "alone", "actual_segment", 0, 10),
        ("alone_later", "alone", "actual_segment", 10, 60),
        ("reunion_first", "reunion", "door_open", 0, 15),
        ("reunion_second", "reunion", "door_open", 15, 30),
        ("stranger_gate", "stranger", "actual_segment", 0, 10),
        ("stranger_enter", "stranger", "actual_segment", 10, 12),
        ("stranger_approach", "stranger", "actual_segment", 12, 16),
        ("stranger_call", "stranger", "name_called", 0, 4),
        ("stranger_contact_plan", "stranger", "actual_segment", 20, 23),
        ("stranger_wait", "stranger", "actual_segment", 23, 25),
        ("stranger_no_contact", "stranger", "actual_segment", 20, 25),
        ("stranger_exit", "stranger", "exit_first_step", 0, None),
        ("floor_shake", "entry", "floor_first_step", 0, 5),
        ("separation_shake", "alone", "owner_left", 0, 5),
    )
    for key, segment, anchor, start, end in timed:
        windows.append({"window_id": key, "segments": [segment], "anchor": anchor,
                        "start_seconds": start, "end_seconds": end,
                        "actual_event_required": anchor != "actual_segment", "source": src})
    for key, segment, anchor in (
            ("reunion_contact", "reunion", "actual_contact"),
            ("stranger_contact", "stranger", "actual_contact"),
            ("entry_object", "entry", "object_one_meter_until_stop_end"),
            ("exit_object", "exit", "object_one_meter_until_stop_end"),
            ("entry_leash", "entry", "entry_movement_excluding_object_stop"),
            ("exit_leash", "exit", "exit_movement_excluding_object_stop"),
            ("walk_start", "walk", "before_first_step"),
            *((f"walk_phase_{number}", "walk", f"actual_phase_{number}") for number in range(1, 7))):
        windows.append({"window_id": key, "segments": [segment], "anchor": anchor,
                        "actual_event_required": True, "source": src})
    windows.append({"window_id": "before_separation", "segments": ["baseline", "alone"],
                    "anchor": "departure_transition_before_owner_exit", "actual_event_required": True,
                    "source": source(root, WORKBOOK, "3_보호자행동!B10")})
    return ProtocolV3.model_validate_json(json.dumps({
        "segment_order": [segment for segment, _, _ in SEGMENTS], "windows": windows,
        "pending_policies": ["Q01", "Q02", "Q03", "Q05", "Q08"],
        "procedure_document": docx_paragraphs(root, PROCEDURE),
        "development_document": docx_paragraphs(root, "00_큐브_회신및개발반영기준_20260929.docx"),
    }, ensure_ascii=False))


def item_windows(prefix: str, row: int, label: str) -> tuple[str, ...]:
    segments = {label: segment for segment, label, _ in SEGMENTS}
    special = {
        "바8": ("floor_shake",), "바10": ("separation_shake",),
        "바42": ("reunion_contact",), "바43": ("reunion_contact",),
        "바44": ("entry_object",), "바45": ("stranger_contact",),
        "개9": ("alone_initial",), "개10": ("alone_later",),
        "개12": ("baseline", "alone"), "개13": ("stranger_gate",),
        "개15": ("stranger_call",), "개17": ("reunion_first",),
        "개18": ("reunion_second",), "개19": ("reunion_contact",),
        "개30": ("entry_object",), "개32": ("entry_object",), "개34": ("exit_object",),
        "개37": ("walk_start",), "개44": ("reunion_contact",),
        "개45": ("entry_leash",), "개46": ("exit_leash",),
        "개51": ("stranger_contact",), "개52": ("stranger_contact",),
        "개54": ("stranger_exit",), "개55": ("stranger_approach",),
        "개60": ("alone_initial", "alone_later", "alone"),
        "보5": ("entry_leash", "exit_leash"), "보6": ("entry_leash", "exit_leash"),
        "보9": ("entry_leash", "exit_leash"), "보10": ("before_separation",),
        "보11": ("reunion_second",), "보14": ("entry_object", "exit_object"),
        "보22": ("reunion_second",), "보39": ("reunion_second",),
        "보25": ("reunion_second",), "보26": ("stranger",), "보27": ("stranger",),
    }
    code = f"{prefix}{row}"
    if code in special:
        return special[code]
    if prefix == "개" and 38 <= row <= 43:
        return (f"walk_phase_{row - 37}",)
    if label in segments:
        return (segments[label],)
    if label == "기록":
        return tuple(segment for segment, _, _ in SEGMENTS)
    if code in UNUSED_CODES:
        return ()
    raise ValueError(f"source window mapping must be explicit: {code} ({label})")


def read_behavior_v3(root: Path, book) -> BehaviorCatalogV3:
    items = []
    for prefix, name, rows in SHEET_ROWS:
        sheet = book[name]
        for row in rows:
            code = f"{prefix}{row}"
            text = sheet[f"B{row}"].value
            if not isinstance(text, str) or not text.strip():
                raise ValueError(f"missing original item text: {code}")
            usage = "unused" if code in UNUSED_CODES else "automatic" if code in AUTO_CODES else "direct"
            checks = [check for check in sheet.data_validations.dataValidation if f"J{row}" in check.sqref]
            if len(checks) > 1:
                raise ValueError(f"ambiguous input validation: {code}")
            check = checks[0] if checks else None
            value_type = usage if usage != "direct" else (
                "memo" if code in ("개59", "보25", "보26", "보27") else
                "seconds" if code == "개32" else
                "count" if prefix == "바" and row in (8, 10, *range(14, 38)) or code == "개21" else "category")
            limits = None
            if check:
                if check.type not in ("whole", "decimal") or check.operator not in (None, "between", "greaterThanOrEqual"):
                    raise ValueError(f"unsupported source validation: {code}")
                limits = {"kind": check.type, "operator": check.operator or "between",
                          "minimum": float(check.formula1),
                          "maximum": float(check.formula2) if check.formula2 is not None else None,
                          "source": source(root, WORKBOOK, f"{name}!{check.sqref}")}
            labels, allowed = [], []
            instruction = sheet[f"C{row}"].value
            if value_type == "category":
                if limits is None or check.type != "whole" or limits["maximum"] is None:
                    raise ValueError(f"category requires explicit whole range: {code}")
                allowed = list(range(int(limits["minimum"]), int(limits["maximum"]) + 1))
                # −1/0/+1 request expression occupies D:F; headers do not apply to other scales.
                columns = "DEF" if code == "보5" else "CDEFG"[:len(allowed)]
                if prefix == "개" and 37 <= row <= 43:
                    parts = re.split(r"\s*/\s*", instruction, maxsplit=3)
                    texts = parts[:3] + [parts[3].split(". 사람", 1)[0]]
                    cells = [f"C{row}"] * 4
                else:
                    texts = [sheet[f"{column}{row}"].value for column in columns]
                    cells = [f"{column}{row}" for column in columns]
                for value, label, cell in zip(allowed, texts, cells, strict=True):
                    if not isinstance(label, str) or not label.strip():
                        raise ValueError(f"missing category label at {name}!{cell}")
                    labels.append({"value": value, "text": label,
                                   "source": source(root, WORKBOOK, f"{name}!{cell}")})
                if not (prefix == "개" and 37 <= row <= 43):
                    instruction = None
            opportunity = {"보22": ("보38",), "보39": ("보38",), "보24": ("보40",)}.get(code, ())
            items.append({"code": code, "sheet": name, "row": row, "text": text,
                          "segment_label": sheet[f"A{row}"].value,
                          "domain_label": sheet[f"H{row}"].value, "axis_label": sheet[f"I{row}"].value,
                          "usage": usage, "value_type": value_type, "labels": labels,
                          "allowed_values": allowed, "input_range": limits, "instruction": instruction,
                          "representative_rule": text + ("\n" + instruction if isinstance(instruction, str) else ""),
                          "windows": item_windows(prefix, row, sheet[f"A{row}"].value),
                          "opportunity_codes": opportunity, "source": source(root, WORKBOOK, f"{name}!A{row}:I{row}")})
    return BehaviorCatalogV3.model_validate_json(json.dumps({"source": source(root, WORKBOOK, "117 original rows"),
                                                           "items": items}, ensure_ascii=False))


def read_survey_v3(root: Path, mapping: dict) -> SurveyCatalogV3:
    with ZipFile(root / SURVEY) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
    questions = {}
    for table_index, table in enumerate(tree.findall(".//w:tbl", NS), 1):
        printed_scale = None
        for row_index, row in enumerate(table.findall("./w:tr", NS), 1):
            cells = ["".join(node.text or "" for node in cell.findall(".//w:t", NS))
                     for cell in row.findall("./w:tc", NS)]
            if len(cells) == 7 and cells[0] == "#" and cells[1].startswith(tuple(f"{section}." for section in "ABCDE")):
                printed_scale = cells[2:]
            if cells and cells[0].isdigit() and 1 <= int(cells[0]) <= 28:
                number = int(cells[0])
                if number in questions:
                    raise ValueError("duplicate printed survey question")
                expected_symbols = ["⓪", "①", "②", "③", "④"] if 10 <= number <= 14 else ["①", "②", "③", "④", "⑤"]
                if printed_scale is None or cells[2:] != expected_symbols:
                    raise ValueError(f"printed scale differs for question {number}")
                questions[number] = (cells[1], printed_scale, f"word/document.xml table {table_index} row {row_index}")
    if set(questions) != set(range(1, 29)):
        raise ValueError("printed questionnaire must contain 28 questions")
    items = []
    for entry in mapping["questions"]:
        number = entry["number"]
        text, printed_scale, location = questions[number]
        if text != entry["text"]:
            raise ValueError(f"DOCX and JSON questionnaire wording differs: {number}")
        labels = ("전혀없다", "약간", "중간정도", "꽤", "심하다") if 10 <= number <= 14 else (
            "전혀아니다", "아니다", "보통", "그렇다", "매우그렇다")
        if tuple(printed_scale) != labels:
            raise ValueError(f"printed response labels differ for question {number}")
        items.append({"item_id": f"s{number:02d}", "external_id": entry["id"], "number": number,
                      "text": text, "print_section": "A" if number <= 9 else "B" if number <= 14 else "C" if number <= 21 else "D" if number <= 25 else "E",
                      "report_domain": entry["group"], "standalone": number == 25, "allowed_values": entry["valid_values"],
                      "reverse_scored": entry["reverse_scored"], "source": source(root, SURVEY, location),
                      "labels": [{"value": value, "text": label, "source": source(root, SURVEY, location)}
                                 for value, label in zip(entry["valid_values"], labels, strict=True)]})
    return SurveyCatalogV3.model_validate_json(json.dumps({"source": source(root, SURVEY, "28 printed questions"),
                                                         "items": items}, ensure_ascii=False))


def expand_questions(values: list[str]) -> tuple[list[str], bool]:
    ids, related = [], False
    for value in values:
        match = re.fullmatch(r"Q(\d{2})(?:-Q(\d{2}))?", value)
        if match:
            first = int(match[1]); last = int(match[2] or match[1])
            if not 1 <= first <= last <= 28:
                raise ValueError("question range outside survey")
            ids.extend(f"s{number:02d}" for number in range(first, last + 1))
        elif value == "필요한 관련 문항":
            related = True
        else:
            raise ValueError(f"unresolved survey mapping expression: {value}")
    return ids, related


def normalize_mapping(root: Path, filename: str, raw: dict, catalog: BehaviorCatalogV3,
                      kind: str) -> MappingV3:
    items = {item.code: item for item in catalog.items}
    entries = []
    source_entries = raw["results"] if kind == "results" else raw["questions"]
    for entry in source_entries:
        questions, related = expand_questions(entry.get("survey", [entry["id"]] if kind == "survey_behavior" else []))
        codes = entry.get("behavior", entry.get("behavior_codes", []))
        bases = []
        if codes == ["각 영역의 확정 근거"]:
            codes = []
            bases = ["owner_type", "attachment_type", "stranger_response", "nonsocial_fear",
                     "walk_proximity", "body_change", "environment_change", "dog_tendency", "voice"]
        if related:
            if entry["id"] != "summary_advice":
                raise ValueError("related question expression lacks an explicit result basis")
            questions = sorted({question for result in raw["results"] if result["id"] in bases
                                for question in expand_questions(result["survey"])[0]})
        windows = list(dict.fromkeys(window for code in codes for window in items[code].windows))
        entries.append({"entry_id": entry["id"], "survey_ids": questions, "behavior_codes": codes,
                        "windows": windows, "result_basis_ids": bases, "original": entry,
                        "source": source(root, filename, f"{kind}/{entry['id']}")})
    return MappingV3.model_validate_json(json.dumps({"kind": kind, "entries": entries}, ensure_ascii=False))


def scoring_rules_v3(root: Path, book) -> ScoringRulesV3:
    formulas = []
    # First slot is a formula template, not a participant limit or imported response.
    for sheet_name, rows, columns in (
            ("2_개행동", range(5, 61), ("J",)),
            ("여러쌍비교", (6,), tuple(cell.column_letter for cell in book["여러쌍비교"][6])),
            ("3_보호자행동", (51,), ("P", "Q", "R", "S", "T", "U"))):
        sheet = book[sheet_name]
        for row in rows:
            for column in columns:
                cell = sheet[f"{column}{row}"]
                if cell.data_type == "f":
                    formulas.append({"formula": cell.value,
                                     "source": source(root, WORKBOOK, f"{sheet_name}!{cell.coordinate}")})
    points = []
    sheet = book["3_보호자행동"]
    for row in range(84, 109):
        points.append({"code": f"보{sheet[f'B{row}'].value}", "raw_value": sheet[f"C{row}"].value,
                       "points": [sheet[f"{column}{row}"].value for column in "DEF"],
                       "scene": sheet[f"G{row}"].value,
                       "source": source(root, WORKBOOK, f"3_보호자행동!B{row}:G{row}")})
    # Added approach-response table is separate from the original six-item table.
    for row in range(147, 151):
        points.append({"code": "보39", "raw_value": sheet[f"C{row}"].value,
                       "points": [sheet[f"{column}{row}"].value for column in "DEF"],
                       "scene": "재회", "source": source(root, WORKBOOK, f"3_보호자행동!B{row}:G{row}")})
    definitions = []
    for name, rows, columns in (("0_안내", (37, 52), ("A",)),
                                ("6_행사전문가의견", range(1, book["6_행사전문가의견"].max_row + 1), ("A", "B", "C", "D", "E", "F", "G")),
                                ("3_보호자행동", range(29, 38), ("A",))):
        for row in rows:
            for column in columns:
                cell = book[name][f"{column}{row}"]
                if cell.value is not None:
                    definitions.append({"text": cell.value, "source": source(root, WORKBOOK, f"{name}!{cell.coordinate}")})
    data = {"version": "scoring-20260929-v3",
            "pending_policies": ["Q01", "Q02", "Q03", "Q04", "Q05", "Q06"],
            "owner_type_thresholds": {"minimum_items": sheet["C77"].value, "minimum_scenes": sheet["C78"].value,
                                      "dominant_ratio": sheet["C79"].value, "ratio_gap": sheet["C80"].value,
                                      "source": source(root, WORKBOOK, "3_보호자행동!B77:C80"), "status": "provisional_v0.2"},
            "owner_type_points": points, "formula_templates": formulas,
            "definitions": definitions, "decision_document": docx_paragraphs(root, DECISIONS),
            "comparison_document": docx_paragraphs(root, "기준자료/03_설문_행동_비교_대응표_20260929.docx"),
            "interpretation_document": docx_paragraphs(root, "기준자료/10_무엇을_보려고_하는가_20260929.docx"),
            "separation_combinations": [{"initial": book["성향문장"][f"A{row}"].value,
                                         "later": book["성향문장"][f"B{row}"].value,
                                         "text": book["성향문장"][f"C{row}"].value,
                                         "source": source(root, WORKBOOK, f"성향문장!A{row}:C{row}")}
                                        for row in range(65, 90)]}
    return ScoringRulesV3.model_validate_json(json.dumps(data, ensure_ascii=False))


def extract_v3(root: Path) -> dict[str, str]:
    manifest = verify_attachments(root)
    survey_raw = json.loads((root / "03_설문28문항_대응표.json").read_text(encoding="utf-8"))
    results_raw = json.loads((root / "02_결과항목_연결표.json").read_text(encoding="utf-8"))
    book = load_workbook(root / WORKBOOK, data_only=False)
    try:
        behavior = read_behavior_v3(root, book)
        survey = read_survey_v3(root, survey_raw)
        protocol = protocol_v3(root)
        results = normalize_mapping(root, "02_결과항목_연결표.json", results_raw, behavior, "results")
        mapping = normalize_mapping(root, "03_설문28문항_대응표.json", survey_raw, behavior, "survey_behavior")
        validate_references_v3(behavior, survey, protocol, (results, mapping))
        rules = scoring_rules_v3(root, book)
    finally:
        book.close()
    if verify_attachments(root) != manifest:
        raise ValueError("source attachments changed during extraction")
    objects = {"catalogs/behavior-v3.json": behavior, "catalogs/survey-v3.json": survey,
               "rules/protocol-v3.json": protocol, "mappings/results-v3.json": results,
               "mappings/survey-behavior-v3.json": mapping, "rules/scoring-v3.json": rules}
    content = {name: model.model_dump_json(indent=2) + "\n" for name, model in objects.items()}
    return content
