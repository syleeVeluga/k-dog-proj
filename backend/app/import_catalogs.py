"""Read-only extraction of customer source workbooks into versioned catalogs; never imports responses."""

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from .domain.catalog import (
    ALL_SEGMENTS_LABEL, BehaviorCatalog, PHASE_COUNT_RANGE, SEGMENTS, SHEET_GROUPS,
)
from .domain.contracts import SurveyCatalog

ROOT = Path(__file__).resolve().parents[2]
CUSTOMER_DIR = ROOT / "docs/최종 고객 문서"
BEHAVIOR_V2_FILE = "03_행동_채점표_42항목_20260913.xlsx"
BEHAVIOR_V2_VERSION = "catalog-20260913-v2"
SURVEY_FILE = "큐브_통합설문_최종양식_20260904.xlsx"
HEADERS = {"A": "구간", "B": "관찰 항목", "C": "1", "D": "2", "E": "3", "F": "4", "G": "5",
           "H": "영역", "I": "척도·축", "AD": "영역코드", "AY": "척도"}
SEGMENT_BY_LABEL = {label: segment for segment, label in SEGMENTS}
# Reference items have no data-type column; the B-column marker decides (계획 R1).
VALUE_TYPE_MARKERS = (("※자동 계산", "auto_ratio"), ("※0~6", "phase_count"), ("※횟수", "count"))


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def literal(cell):
    if cell.data_type == "f":
        raise ValueError(f"expected a literal value at {cell.parent.title}!{cell.coordinate}")
    return cell.value


def read_behavior_v2(path: Path) -> BehaviorCatalog:
    source_hash = file_hash(path)
    book = load_workbook(path, data_only=False)
    items = []
    try:
        for prefix, name, count in SHEET_GROUPS:
            sheet = book[name]
            for column, expected in HEADERS.items():
                if sheet[f"{column}4"].value != expected:
                    raise ValueError(f"unexpected header {name}!{column}4")
            for number in range(1, count + 1):
                row = number + 4
                text = literal(sheet[f"B{row}"])
                if not isinstance(text, str):
                    raise ValueError(f"item text missing at {name}!B{row}")
                segment_label = literal(sheet[f"A{row}"])
                if segment_label != ALL_SEGMENTS_LABEL and segment_label not in SEGMENT_BY_LABEL:
                    raise ValueError(f"unknown segment label at {name}!A{row}")
                axis_label = literal(sheet[f"I{row}"])
                if not isinstance(axis_label, str):
                    raise ValueError(f"axis label missing at {name}!I{row}")
                domain = literal(sheet[f"AD{row}"])
                scale = literal(sheet[f"AY{row}"])
                value_type = next((kind for marker, kind in VALUE_TYPE_MARKERS if marker in text), "scale")
                cells = [sheet[f"{column}{row}"] for column in "CDEFG"]
                labels = [{"score": score, "text": literal(cell), "source_cell": cell.coordinate}
                          for score, cell in enumerate(cells, 1) if cell.value is not None]
                if value_type == "scale":
                    allowed, note = [label["score"] for label in labels], None
                else:
                    if len(labels) != 1 or labels[0]["score"] != 1:
                        raise ValueError(f"count/auto item needs a single instruction in C at {name}!C{row}")
                    allowed, note, labels = (list(PHASE_COUNT_RANGE) if value_type == "phase_count" else []), labels[0]["text"], []
                items.append({
                    "item_id": f"{prefix}-{number:02d}", "sheet": name, "source_row": row,
                    "segment_label": segment_label,
                    "segment": None if segment_label == ALL_SEGMENTS_LABEL else SEGMENT_BY_LABEL[segment_label],
                    "text": text, "domain_label": literal(sheet[f"H{row}"]), "domain": domain,
                    "axis_label": axis_label, "axis": axis_label.split("·", 1)[1] if "·" in axis_label else None,
                    "scale": scale, "value_type": value_type, "labels": labels,
                    "allowed_scores": allowed, "note": note,
                })
            if sheet[f"B{5 + count}"].value is not None:
                raise ValueError(f"{name} has more item rows than the agreed {count}")
    finally:
        book.close()
    if file_hash(path) != source_hash:
        raise ValueError("source workbook changed during extraction")
    return BehaviorCatalog.model_validate_json(json.dumps({
        "version": BEHAVIOR_V2_VERSION, "source_filename": path.name,
        "source_sha256": source_hash, "provenance": "excel_verified", "items": items,
    }, ensure_ascii=False))


def read_survey(path: Path) -> SurveyCatalog:
    source_hash = file_hash(path)
    book = load_workbook(path, read_only=True, data_only=False)
    items = []
    try:
        sheet = book["데이터입력"]
        instructions = sheet["A2"].value
        if not isinstance(instructions, str) or "1~5" not in instructions:
            raise ValueError("survey response scale needs review")
        for number, row in enumerate(sheet.iter_rows(min_row=6, max_row=35, max_col=5), 1):
            if row[0].value != number or any(cell.data_type == "f" for cell in row):
                raise ValueError("survey numbering or literal content changed")
            items.append({
                "item_id": f"q{number:02d}", "text": row[3].value,
                "domain_label": row[1].value, "source_layer": row[2].value,
                "scoring_note": row[4].value, "source_sheet": sheet.title,
                "source_row": number + 5,
            })
    finally:
        book.close()
    if file_hash(path) != source_hash:
        raise ValueError("source workbook changed during extraction")
    return SurveyCatalog.model_validate_json(json.dumps({
        "version": "catalog-20260904-v1", "source_filename": path.name,
        "source_sha256": source_hash, "provenance": "excel_verified",
        "response_instructions": instructions, "allowed_responses": [1, 2, 3, 4, 5],
        "items": items,
    }, ensure_ascii=False))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--customer-dir", type=Path, default=CUSTOMER_DIR, help="folder holding the 42-item customer workbooks")
    parser.add_argument("--source-dir", type=Path, default=ROOT / "resources/source", help="folder holding the superseded 55-item workbooks")
    parser.add_argument("--check", action="store_true", help="compare existing JSON without writing")
    args = parser.parse_args()
    catalogs = {
        "behavior-v2.json": read_behavior_v2(args.customer_dir / BEHAVIOR_V2_FILE),
        "survey-v1.json": read_survey(args.source_dir / SURVEY_FILE),
    }
    for name, catalog in catalogs.items():
        destination = ROOT / "resources" / "catalogs" / name
        content = catalog.model_dump_json(indent=2) + "\n"
        if args.check:
            if not destination.exists() or destination.read_text(encoding="utf-8") != content:
                raise ValueError(f"catalog is absent or differs from source: {name}")
        else:
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_text(content, encoding="utf-8", newline="\n")
        print(f"{name}: {len(catalog.items)} items; source SHA-256 {catalog.source_sha256}")


if __name__ == "__main__":
    main()
