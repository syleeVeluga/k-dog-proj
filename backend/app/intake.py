"""Participant and 28-item survey intake, shared by manual entry and standard imports."""

import csv
from io import BytesIO, StringIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.domain.catalog import SURVEY_IDS, SurveyCatalog
from app.domain.contracts import SurveyAnswers
from app.domain.validation import validate_survey_answers
from app.input_models import CaseCreate, ImportPreview, ImportRow, Manifest, Session, SurveyEdit
from app.storage import now, uid

# 04 설문지: 7~9번은 「해당 없음」 칸이 있다. 가져오기 파일에서는 이 문자열로 표시한다.
NOT_APPLICABLE_TOKENS = ("NA", "na", "N/A", "해당없음", "해당 없음")


def new_session(survey_version: str, note="") -> Session:
    return Session(session_id=uid(), note=note, survey_version=survey_version,
                   survey=dict.fromkeys(SURVEY_IDS), survey_not_applicable=[], videos=[])


def create_case(store, db, value: CaseCreate, actor: str, version: str):
    session = new_session(version)
    manifest = Manifest(case_id=uid(), event_id=value.event_id,
                        participant_id=value.participant_id, input_revision=1,
                        selected_session_id=session.session_id, sessions=[session])
    key, digest = store.write_manifest(manifest)
    db.execute(
        "INSERT INTO cases(case_id,event_id,participant_id,dog_name,reservation_at,sequence_no,consent_confirmed,guardian_name,"
        "dog_profile_json,input_revision,selected_session_id,manifest_ref,manifest_hash,created_at,updated_at) "
        "VALUES (?,?,?,?,?,?,?,?,?,1,?,?,?,?,?)",
        (manifest.case_id, value.event_id, value.participant_id, value.dog_name, value.reservation_at,
         value.sequence_no, int(value.consent_confirmed), value.guardian_name, value.dog.model_dump_json(),
         session.session_id, key, digest, now(), now()),
    )
    store.audit(db, actor, manifest.case_id, "case.create", {"revision": 1})
    return manifest.case_id


def selected_session(manifest, session_id):
    if manifest.selected_session_id != session_id:
        raise HTTPException(409, "선택된 촬영 세션이 변경되었습니다. 다시 조회하세요.")
    return next(item for item in manifest.sessions if item.session_id == session_id)


def save_survey(store, db, case_id, value: SurveyEdit, actor, catalog: SurveyCatalog):
    if value.survey_version != catalog.version:
        raise HTTPException(422, "확정 설문 버전과 일치하지 않습니다.")
    try:
        validate_survey_answers(SurveyAnswers(answers=value.answers, not_applicable=tuple(value.not_applicable)), catalog)
    except ValueError as exc:
        raise HTTPException(422, "「해당 없음」은 7~9번 문항에만 표시할 수 있습니다.") from exc
    row = store.case(db, case_id, expected=value.expected_revision)
    manifest = store.manifest(row)
    session = selected_session(manifest, value.session_id)
    not_applicable = sorted(value.not_applicable)
    if session.survey == value.answers and session.survey_not_applicable == not_applicable:
        return  # Identical re-import is idempotent.
    session.survey = value.answers
    session.survey_not_applicable = not_applicable
    store.save(db, row, manifest, actor, "survey.update")


PROFILE_COLUMNS = ("dog_breed", "dog_sex", "dog_age_years", "dog_size", "years_together", "adoption_route")
TRUE_TOKENS = ("1", "Y", "y", "예", "동의", "true", "True", "O")


def headers(kind):
    if kind == "participants":
        return ["event_id", "participant_id", "dog_name", "reservation_at", "sequence_no", "consent_confirmed", "guardian_name", *PROFILE_COLUMNS]
    return ["event_id", "participant_id", "survey_version", *SURVEY_IDS]


def required_columns(kind):
    """A participants file may omit the optional 접수 columns; a survey file must carry every item."""
    return {"event_id", "participant_id", "dog_name"} if kind == "participants" else set(headers(kind))


def participant_row(value):
    """Map spreadsheet cells to CaseCreate: blanks mean 「미기재」, never a guessed value."""
    def text(key):
        raw = value.get(key)
        return "" if raw is None else str(raw).strip()

    def integer(key):
        raw = value.get(key)
        if raw is None or raw == "":
            return None
        if type(raw) is int or (type(raw) is str and raw.strip().isdigit()):
            return int(raw)
        raise ValueError(f"{key}: 정수만 허용됩니다.")

    return CaseCreate.model_validate({
        "event_id": value["event_id"], "participant_id": value["participant_id"], "dog_name": value["dog_name"],
        "reservation_at": text("reservation_at"), "sequence_no": integer("sequence_no"),
        "consent_confirmed": text("consent_confirmed") in TRUE_TOKENS, "guardian_name": text("guardian_name"),
        "dog": {"breed": text("dog_breed"), "sex": text("dog_sex") or "미기재", "age_years": integer("dog_age_years"),
                "size": text("dog_size") or "미기재", "years_together": text("years_together"), "adoption_route": text("adoption_route") or "미기재"},
    })


def template(kind, format):
    columns = headers(kind)
    if format == "csv":
        output = StringIO(newline="")
        csv.writer(output).writerow(columns)
        return output.getvalue().encode("utf-8-sig")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "입력"
    sheet.append(columns)
    # IDs must remain text, including leading zeroes.
    for row in sheet.iter_rows(min_row=2, max_row=101, max_col=len(columns)):
        for cell in row:
            cell.number_format = "@"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def read_rows(data: bytes, format: str, sheet_name=None):
    if format == "csv":
        try:
            return list(csv.reader(StringIO(data.decode("utf-8-sig")), strict=True))
        except (UnicodeError, csv.Error) as exc:
            raise HTTPException(422, "UTF-8 CSV 형식을 확인하세요.") from exc
    try:
        # Reject oversized expanded XML before parsing a compressed workbook.
        from zipfile import ZipFile
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 32 * 1024 * 1024:
                raise ValueError("expanded workbook too large")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            sheet = workbook[sheet_name] if sheet_name else workbook.active
            if (sheet.max_row or 0) > 10001 or (sheet.max_column or 0) > 100:
                raise ValueError("too many cells")
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    except Exception as exc:
        raise HTTPException(422, "표준 XLSX 파일을 확인하세요. 수식 대신 원응답을 입력하세요.") from exc


def mapped_rows(rows, kind, version, mapping):
    if not mapping.columns:
        return rows
    expected = headers(kind)
    if set(mapping.columns) - set(expected):
        raise HTTPException(422, "알 수 없는 표준 열 연결입니다.")
    source = list(mapping.columns.values())
    if len(source) != len(set(source)) or any(rows[0].count(name) != 1 for name in source):
        raise HTTPException(422, "각 원본 열을 중복 없이 연결하세요.")
    if not required_columns(kind) - {"survey_version"} <= set(mapping.columns):
        raise HTTPException(422, "모든 필수 열과 28문항을 연결하세요.")
    return [expected, *[[row[rows[0].index(mapping.columns[name])] if name in mapping.columns and rows[0].index(mapping.columns[name]) < len(row)
                        else version if name == "survey_version" else "" for name in expected] for row in rows[1:]]]


def parse_answer(question, raw, catalog):
    """Return (answer, not_applicable) for one survey cell; blanks stay missing, 해당 없음 only where the form offers it."""
    if raw is None or raw == "":
        return None, False
    if isinstance(raw, str) and raw.strip() in NOT_APPLICABLE_TOKENS:
        if not next(item for item in catalog.items if item.item_id == question).allows_not_applicable:
            raise ValueError(f"{question}: 「해당 없음」은 7~9번 문항에만 허용됩니다.")
        return None, True
    if (type(raw) is int and 1 <= raw <= 5) or (type(raw) is str and raw in ("1", "2", "3", "4", "5")):
        return int(raw), False
    raise ValueError(f"{question}: 1~5, 빈칸, 또는 7~9번의 NA만 허용됩니다.")


def preview(store, db, data, kind, format, catalog: SurveyCatalog, mapping=None):
    version = catalog.version
    rows = read_rows(data, format, mapping.sheet if mapping else None)
    if rows and mapping:
        rows = mapped_rows(rows, kind, version, mapping)
    if not rows or len(rows) > 10001:
        raise HTTPException(422, "헤더와 최대 10,000개 입력 행이 필요합니다.")
    columns = rows[0]
    if len(columns) != len(set(columns)) or not set(columns) <= set(headers(kind)) or not required_columns(kind) <= set(columns):
        raise HTTPException(422, "표준 양식의 열을 사용하세요. 열 순서 변경은 허용되며 참가자 양식의 접수 열은 생략할 수 있습니다.")
    result = ImportPreview(rows=[], errors=[])
    seen = set()
    for number, values in enumerate(rows[1:], 2):
        location = f"{number}행"
        if all(value is None or value == "" for value in values):
            continue
        try:
            if len(values) != len(columns):
                raise ValueError("열 수가 헤더와 다릅니다.")
            value = dict(zip(columns, values))
            for key in ("event_id", "participant_id"):
                if not isinstance(value[key], str) or not value[key].strip():
                    raise ValueError(f"{key}: 텍스트 ID가 필요합니다. Excel 숫자 ID는 자동 변환하지 않습니다.")
            pair = value["event_id"], value["participant_id"]
            if pair in seen:
                raise ValueError("파일 안에 중복 참가자 ID가 있습니다.")
            seen.add(pair)
            existing = db.execute("SELECT * FROM cases WHERE event_id=? AND participant_id=?", pair).fetchone()
            if kind == "participants":
                if existing:
                    raise ValueError("이미 등록된 참가자 ID입니다.")
                result.rows.append(ImportRow(row_number=number, source_location=location, event_id=pair[0], participant_id=pair[1],
                                             participant=participant_row(value)))
            else:
                if existing is None or existing["deletion_requested"]:
                    raise ValueError("등록된 활성 참가자에 연결할 수 없습니다.")
                answers, not_applicable = {}, []
                for question in SURVEY_IDS:
                    answers[question], flagged = parse_answer(question, value[question], catalog)
                    if flagged:
                        not_applicable.append(question)
                if value["survey_version"] != version:
                    raise ValueError("survey_version: 확정 설문 버전이 아닙니다.")
                survey = SurveyEdit(expected_revision=existing["input_revision"], session_id=existing["selected_session_id"],
                                    survey_version=version, answers=answers, not_applicable=not_applicable)
                manifest = store.manifest(existing)
                index = next(i for i, s in enumerate(manifest.sessions) if s.session_id == existing["selected_session_id"])
                current = manifest.sessions[index]
                changed = [q for q in SURVEY_IDS if current.survey[q] != answers[q] or (q in current.survey_not_applicable) != (q in not_applicable)]
                result.rows.append(ImportRow(row_number=number, source_location=location, event_id=pair[0], participant_id=pair[1],
                    case_id=existing["case_id"], survey=survey, session_label=f"{index + 1}차 촬영", changed_questions=changed))
        except (ValueError, ValidationError) as exc:
            if isinstance(exc, ValidationError):
                explanation = "; ".join(".".join(map(str, e["loc"])) + ": " + e["msg"] for e in exc.errors())
            else:
                explanation = str(exc)
            result.errors.append(f"{location}: {explanation}")
    return result
