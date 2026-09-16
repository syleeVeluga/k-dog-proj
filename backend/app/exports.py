"""S7 frozen batches and real PDF/XLSX/CSV rendering, with no provider calls."""

import base64
import csv
from datetime import datetime, timezone
import hashlib
from io import BytesIO, StringIO
import json
import os
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Table, TableStyle, PageBreak, Spacer

from app.reporting import NOTICE, PRESENTATION, read_saved, run_row
from app.storage import REPO_ROOT, encode, uid


def guard_snapshot(store, db, snapshot):
    for member in snapshot["members"]:
        case = store.case(db, member["case_id"])
        if member["run_id"]:
            run_row(store, db, case["case_id"], member["run_id"])


def load_snapshot(store, db, export_id):
    row = db.execute("SELECT detail_json FROM changes WHERE target=? AND action='export.snapshot'", (export_id,)).fetchone()
    if row is None:
        raise HTTPException(404, "내보내기를 찾을 수 없습니다.")
    snapshot = read_saved(store, json.loads(row[0]))
    if snapshot["export_id"] != export_id:
        raise HTTPException(409, "내보내기 참조가 일치하지 않습니다.")
    guard_snapshot(store, db, snapshot)
    return snapshot


def export_view(snapshot, status="snapshot"):
    members = [{k: member.get(k) for k in ("case_id", "event_id", "participant_id", "dog_name", "input_revision", "revision", "run_id", "status", "explanation_status")}
               for member in snapshot["members"]]
    return {"export_id": snapshot["export_id"], "format": snapshot["format"], "created_at": snapshot["created_at"],
            "actor": snapshot["actor"], "count": len(snapshot["members"]), "status": status,
            "preview_hash": snapshot.get("preview_hash", ""), "members": members}


def report_coverage(data):
    data = data or {}
    branches = {a["evaluation"]["branch"] for a in data.get("evaluations", [])}
    items = (data.get("scores") or {}).get("items", [])
    total = len(data.get("behavior_items", []))
    scored = sum(i["status"] == "scored" for i in items)
    state = "부분 평가 결과" if branches != {"dog", "owner"} or scored < total else "평가 범위"
    return (f"{state} · 반려견 평가 {'완료' if 'dog' in branches else '미완료'} · 보호자 평가 {'완료' if 'owner' in branches else '미완료'}"
            + (f" · {total}개 항목 중 {scored}개 채점. 미채점은 0점이 아닙니다." if total else " · 채점 결과 없음"))


def tabular(snapshot):
    presentation = snapshot.get("presentation", PRESENTATION)
    topics = presentation["topics"]
    titles = {t["slot"]: t["title"] + " (잠정)" for t in topics}
    tables = {
        "전체요약": [["행사", "참가자 ID", "반려견", "상태", "입력 버전", "수정 버전", "실행 ID", "설명 상태", "생성시각 UTC", "내보내기 ID"]],
        "리포트": [["행사", "참가자 ID", "구성", "설명", "근거 ID"]],
        "영역비교": [["행사", "참가자 ID", "비교 항목 번호", "보호자 설문", "영상 관찰", "상태", "평가 주제 (잠정)"]],
        "행동55항목": [["행사", "참가자 ID", "항목 ID", "항목", "영역", "상태", "선택지", "점수", "방향", "사유", "근거 ID"]],
        "행동집계": [["행사", "참가자 ID", "영역", "평균", "최고", "유효 수", "대상 수"]],
        "설문": [["행사", "참가자 ID", "문항", "원문", "원응답", "환산값", "상태"]],
        "근거": [["행사", "참가자 ID", "근거 ID", "영상 ID", "카메라", "시작초", "종료초", "관찰", "항목 ID"]],
        "수정이력": [["행사", "참가자 ID", "수정 버전", "사용자", "시각 UTC", "종류", "사유", "이전값", "수정값"]],
        "규칙설명": [["구분", "내용"], ["평가 주제 안내", presentation["notice"]], *[[k, v] for k, v in snapshot["rules"].items()]],
    }
    for m in snapshot["members"]:
        identity = [m["event_id"], m["participant_id"]]
        tables["전체요약"].append([*identity, m["dog_name"], m["status"], m["input_revision"], m["revision"], m["run_id"], m["explanation_status"],
                                  datetime.fromisoformat(snapshot["created_at"]).astimezone(timezone.utc).replace(tzinfo=None), snapshot["export_id"]])
        for slot in range(1, 5):
            tables["영역비교"].append([*identity, slot, None, None, "비교 기준 확인 중", titles[slot]])
        report = m["report"]
        parts = [("이번 평가 요약", report["cover"]["text"], report["cover"]["evidence_ids"])] if report else [("이번 평가 요약", "아직 평가 설명이 준비되지 않았습니다.", [])]
        parts += [("평가 범위", report_coverage(m["result"]), []), ("평가 주제 안내", presentation["notice"], [])]
        comments = {d["slot"]: d for d in report["domains"]} if report else {}
        for topic in topics:
            comment = comments.get(topic["slot"], {})
            parts.append((titles[topic["slot"]], topic["description"] + "\n" + comment.get("comment", presentation["pending"]), comment.get("evidence_ids", [])))
        parts += [("관계 스타일 해설", report["cross_type"]["explanation"] if report else presentation["cross_pending"], report["cross_type"]["evidence_ids"] if report else [])]
        parts += [("오늘의 팁", tip["text"], tip["evidence_ids"]) for tip in report["tips"]] if report else [("오늘의 팁", "근거 기반 설명 준비 후 제공됩니다.", [])]
        parts += [("안내", NOTICE, [])]
        tables["리포트"].extend([[*identity, title, text, ", ".join(ids)] for title, text, ids in parts])
        data = m["result"] or {}
        scores = {i["item_id"]: i for i in (data.get("scores") or {}).get("items", [])}
        choices = {i["item_id"]: i for a in data.get("evaluations", []) for i in a["evaluation"]["items"]}
        for item in snapshot["catalog"]["items"]:
            s, choice = scores.get(item["item_id"], {}), choices.get(item["item_id"], {})
            selected = next((o["text"] for o in item["options"] if o["option_id"] == choice.get("selected_option_id")), None)
            tables["행동55항목"].append([*identity, item["item_id"], item["text"], item["domain"], s.get("status", "not_processed"), selected,
                s.get("raw_score"), s.get("direction"), s.get("reason", "평가 미완료"), ", ".join(choice.get("evidence_ids", []))])
        for d in (data.get("scores") or {}).get("domains", []):
            tables["행동집계"].append([*identity, d["domain"], d["mean"], d["maximum"], d["valid_count"], d["target_count"]])
        survey = {q["item_id"]: q for q in (data.get("survey_scores") or {}).get("items", [])}
        for q in snapshot["survey_catalog"]["items"]:
            s = survey.get(q["item_id"], {})
            tables["설문"].append([*identity, q["item_id"], q["text"], s.get("raw", m.get("raw_survey", {}).get(q["item_id"])), s.get("converted"), s.get("status", "not_processed")])
        for e in data.get("evidence", []):
            tables["근거"].append([*identity, e["evidence_id"], e["video_id"], e["camera_id"], e["source_start_sec"], e["source_end_sec"], e["observation"], ", ".join(e["candidate_item_ids"])])
        for h in m["history"]:
            tables["수정이력"].append([*identity, h["revision"], h["actor"], datetime.fromisoformat(h["at"]).astimezone(timezone.utc).replace(tzinfo=None), h["kind"], h["reason"], encode(h.get("before")), encode(h.get("after"))])
    return tables


def workbook(snapshot):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabular(snapshot).items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "C2" if name != "규칙설명" else "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    # Literal strings, including = + - @, must never become formulas.
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.font = Font(name="NanumGothic", size=11, color="253735")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.00"
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="294A46")
                    cell.font = Font(name="NanumGothic", size=11, color="FFFFFF", bold=True)
            ws.row_dimensions[row[0].row].height = 32 if row[0].row == 1 else 64
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = 18 if len(column) == 1 else min(65, max(18, max(len(str(c.value or "")) for c in column) * 1.2))
        if name == "리포트":
            ws.column_dimensions["D"].width = 85
            for i in range(2, ws.max_row + 1):
                ws.row_dimensions[i].height = min(409, max(42, len(str(ws.cell(i, 4).value)) / 45 * 18))
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
    # Four missing comparisons are explicitly blank, never a zero-height score chart.
    ws = wb["영역비교"]
    ws["H1"] = "보호자 설문 / 영상 관찰 비교"
    ws["H2"] = snapshot.get("presentation", PRESENTATION)["notice"]
    ws.column_dimensions["H"].width = 70
    ws["H2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = max(ws.row_dimensions[2].height or 0, 120)
    for index, m in enumerate(snapshot["members"], 1):
        if m["image"]:
            ws = wb.create_sheet(f"대표이미지{index}")
            ws.append([m["event_id"], m["participant_id"], m["dog_name"]])
            for c in ws[1]:
                c.data_type = "s"
                c.number_format = "@"
            ws.append(["원본 영상", m["image"]["video_id"], m["image"]["second"]])
            image = SheetImage(BytesIO(base64.b64decode(m["image"]["data"])))
            scale = min(1, 560 / image.width, 420 / image.height)
            image.width, image.height = image.width * scale, image.height * scale
            ws.add_image(image, "A4")
            for col, width in (("A", 24), ("B", 45), ("C", 24)):
                ws.column_dimensions[col].width = width
            for row in ws:
                for cell in row:
                    cell.font = Font(name="NanumGothic", size=11)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[1].height = 32
            ws.row_dimensions[2].height = 40
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.print_area = "A1:C30"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def pdf(snapshot, member):
    presentation = snapshot.get("presentation", PRESENTATION)
    if "KDog" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("KDog", str(REPO_ROOT / "resources/fonts/NanumGothic-Regular.ttf")))
    if "KDogSymbols" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("KDogSymbols", str(REPO_ROOT / "resources/fonts/NotoSansSymbols.ttf")))
    style = ParagraphStyle("body", fontName="KDog", fontSize=11, leading=18, wordWrap="CJK", spaceAfter=10)
    title_style = ParagraphStyle("title", parent=style, fontSize=23, leading=30, spaceAfter=16)
    heading = ParagraphStyle("heading", parent=style, fontSize=14, leading=21, spaceBefore=16, keepWithNext=True)
    def p(text, selected=style):
        text = escape(str(text)).replace("\n", "<br/>")
        for symbol in "②④":
            text = text.replace(symbol, f'<font name="KDogSymbols">{symbol}</font>')
        return Paragraph(text, selected)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=(595.28, 841.89), rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=46,
                            title="K-DOG 관찰 리포트", author="K-DOG")
    m = member
    story = [p("K-DOG 관찰 리포트", title_style), p(m["dog_name"], heading),
             p(f"{m['event_id']} / 참가자 {m['participant_id']} · 입력 {m['input_revision']} · 수정 {m['revision'] or '-'}"),
             p(f"평가 상태 {m['status']} · 설명 상태 {m['explanation_status']}"),
             p("AI 관찰·평가 기반 자료이며 사람 수정은 아래 이력에 표시됩니다.")]
    if m["image"]:
        img = Image(BytesIO(base64.b64decode(m["image"]["data"])))
        ratio = min(490 / img.imageWidth, 205 / img.imageHeight)
        img.drawWidth, img.drawHeight = img.imageWidth * ratio, img.imageHeight * ratio
        story += [img, p(f"대표 프레임 · 원본 {m['image']['second']:.2f}초")]
    else:
        story.append(p("대표 이미지 미선택"))
    report = m["report"]
    story += [p("이번 평가 요약", heading), p(report["cover"]["text"] if report else "아직 평가 설명이 준비되지 않았습니다."), p(report_coverage(m["result"])),
              p("보호자 설문·영상 관찰 비교", heading), p(presentation["notice"])]
    table = Table([[p("평가 주제 (잠정)"), p("보호자 설문"), p("영상 관찰")],
                   *[[p(t["title"]), p("비교 기준 확인 중"), p("비교 기준 확인 중")] for t in presentation["topics"]]], colWidths=[170, 166, 166], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EEEB")), ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD8D3")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [table, PageBreak(), p("평가 주제별 설명", title_style), p("주제명과 연결 기준은 잠정안입니다. 아래는 저장된 설명과 근거입니다.")]
    evidence = {e["evidence_id"]: e for e in (m["result"] or {}).get("evidence", [])}
    def cite(ids):
        return " / ".join(f"{evidence[i]['camera_id']} {evidence[i]['source_start_sec']:.2f}초: {evidence[i]['observation']}" for i in ids if i in evidence) or "연결 근거 없음"
    comments = {d["slot"]: d for d in report["domains"]} if report else {}
    for topic in presentation["topics"]:
        d = comments.get(topic["slot"], {"comment": presentation["pending"], "evidence_ids": []})
        story += [p(topic["title"] + " (잠정)", heading), p(topic["description"]), p(d["comment"]), p("근거: " + cite(d["evidence_ids"]))]
    story += [p("관계 스타일 해설", heading), p(report["cross_type"]["explanation"] if report else presentation["cross_pending"]), p("오늘의 팁", heading)]
    for tip in report["tips"] if report else [{"text": "근거 기반 설명 준비 후 제공됩니다.", "evidence_ids": []}]:
        story += [p(tip["text"]), p("근거: " + cite(tip["evidence_ids"]))]
    story += [p(NOTICE), Spacer(1, 24), p("점수·검토 기록", heading)]
    for d in ((m["result"] or {}).get("scores") or {}).get("domains", []):
        story.append(p(f"{d['domain']} · 평균 {d['mean'] if d['mean'] is not None else '미산출'} · 최고 {d['maximum'] if d['maximum'] is not None else '미산출'} · 유효 {d['valid_count']}/{d['target_count']}"))
    for h in m["history"]:
        story.append(p(f"수정 {h['revision']} · {h['actor']} · {h['at']} · {h['reason']}"))
    if not m["history"]:
        story.append(p("사람 수정 이력 없음. 교수 검토 완료를 의미하지 않습니다."))
    story += [p("q23/C-2·DOG-12·OWN-14 계산 규칙 보류"), p(f"실행: {m['run_id'] or '미실행'}"), p(f"내보내기: {snapshot['export_id']}"), p(f"생성자: {snapshot['actor']} · {snapshot['created_at']}")]
    def footer(canvas, doc):
        canvas.setFont("KDog", 8)
        canvas.drawString(42, 25, f"K-DOG · {m['participant_id']} · {doc.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render(snapshot):
    if snapshot["format"] == "xlsx":
        return workbook(snapshot), "xlsx"
    if snapshot["format"] == "pdf" and snapshot["individual"]:
        return pdf(snapshot, snapshot["members"][0]), "pdf"
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        if snapshot["format"] == "pdf":
            for m in snapshot["members"]:
                archive.writestr(f"{m['event_id']}/{m['participant_id']}.pdf", pdf(snapshot, m))
        else:
            for name, rows in tabular(snapshot).items():
                stream = StringIO(newline="")
                writer = csv.writer(stream)
                for row in rows:
                    # CSV has no cell typing; prefix executable text for spreadsheet safety.
                    writer.writerow(["'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else v for v in row])
                archive.writestr(name + ".csv", stream.getvalue().encode("utf-8-sig"))
    return output.getvalue(), "zip"


def generate(store, export_id, actor):
    with store.connect() as db:
        snapshot = load_snapshot(store, db, export_id)
        existing = db.execute("SELECT detail_json FROM changes WHERE target=? AND action='export.file' ORDER BY rowid DESC LIMIT 1", (export_id,)).fetchone()
    if existing:
        link = json.loads(existing[0])
        path = store.path(link["ref"])
        if hashlib.sha256(path.read_bytes()).hexdigest() != link["hash"]:
            raise HTTPException(409, "내보내기 파일 해시가 일치하지 않습니다.")
        with store.connect() as db:
            guard_snapshot(store, db, snapshot)
        return path
    raw, extension = render(snapshot)
    # Publish bytes atomically, then adopt only after checking current access again.
    path = store.path(f"exports/{export_id}/result.{extension}")
    temp = path.with_name(uid() + ".tmp")
    try:
        with temp.open("xb") as handle:
            handle.write(raw)
            handle.flush()
            os.fsync(handle.fileno())
        with store.connect(write=True) as db:
            guard_snapshot(store, db, snapshot)
            adopted = db.execute("SELECT detail_json FROM changes WHERE target=? AND action='export.file' LIMIT 1", (export_id,)).fetchone()
            if adopted:
                link = json.loads(adopted[0])
                path = store.path(link["ref"])
                if hashlib.sha256(path.read_bytes()).hexdigest() != link["hash"]:
                    raise HTTPException(409, "내보내기 파일 해시가 일치하지 않습니다.")
            else:
                # An unadopted file after interruption is regenerated from the same snapshot.
                os.replace(temp, path)
                store.audit(db, actor, export_id, "export.file", {"ref": path.relative_to(store.root).as_posix(), "hash": hashlib.sha256(raw).hexdigest()})
    finally:
        temp.unlink(missing_ok=True)
    return path
